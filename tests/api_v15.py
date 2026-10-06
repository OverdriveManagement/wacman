"""V1.5 : relevé des actions, registre des décisions, e-mail du CR, quoi de neuf, revue de stream, bilan de sprint,
import d'un CR d'atelier (Claude simulé), brouillon de faits marquants, slides des decks Program weekly et COPROJ."""
import os as _os
ENV_FILE = _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "apps", "api", ".env.dev")
import sys, json, io, base64, copy, subprocess, urllib.request, http.cookiejar

API = "http://localhost:4000"
ENV = open(ENV_FILE).read()
PWD = ENV.split("BOOTSTRAP_ADMIN_PASSWORD=")[1].split("\n")[0].strip()
DB = ENV.split("DATABASE_URL=")[1].split("\n")[0].strip()
ACC = "la-poste-pstng"
VIEWER = ("lecteur.tour@wacman.test", "Lecteur-Tour-2026-ok!")
fails = []


def check(name, cond, extra=""):
    print(("OK  " if cond else "FAIL"), name, extra)
    if not cond:
        fails.append(name)


def client():
    cj = http.cookiejar.CookieJar()
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))

    def req(method, path, body=None, headers=None):
        data = json.dumps(body).encode() if body is not None else None
        h = dict(headers or {})
        if data is not None:
            h["Content-Type"] = "application/json"
        r = urllib.request.Request(API + path, data=data, method=method, headers=h)
        try:
            with op.open(r) as res:
                t = res.read()
                return res.status, (json.loads(t) if t and "json" in res.headers.get("content-type", "") else t)
        except urllib.error.HTTPError as e:
            t = e.read().decode()
            try:
                return e.code, json.loads(t)
            except Exception:
                return e.code, t

    return req


def login(req, email, pwd):
    s, r = req("POST", "/api/auth/login", {"email": email, "password": pwd})
    req("POST", "/api/auth/verify", {"challengeId": r["challengeId"], "code": r["devCode"]})


req = client()
login(req, "florent@omgt.fr", PWD)
boot = req("GET", f"/api/accounts/{ACC}")[1]
types = {t["name"]: t for t in boot["meetingTypes"]}
streams = boot["streams"]
coproj = next(t for t in boot["meetingTypes"] if "COPROJ" in t["name"].upper())
weekly = next((t for t in boot["meetingTypes"] if "weekly" in t["name"].lower()), boot["meetingTypes"][0])
sprint = next(s for s in boot["sprints"] if s["state"] == "CURRENT")

# ---------------------------------------------------------------- reprise des données (migration 0004)
check("COPROJ : bloc ACTIONS", "ACTIONS" in coproj["blocks"], str(coproj["blocks"]))
strat = [t for t in boot["meetingTypes"] if "TOPICS" in t["blocks"]]
check("séances à sujets : bloc DECISIONS", all("DECISIONS" in t["blocks"] for t in strat), str([t["name"] for t in strat]))
st = coproj.get("settings") or {}
check("COPROJ : objet d'e-mail", st.get("mailSubject") == "WIFIRST / PSTNG : CR COPROJ du {date}", str(st.get("mailSubject")))
check("COPROJ : destinataires", "tuyen.vu-prestataire@laposte.fr" in (st.get("mailTo") or "") and "matthieu.roca@wifirst.fr" in (st.get("mailCc") or ""))
s, acts = req("GET", f"/api/accounts/{ACC}/e/action")
check("actions reprises du COPROJ du 01/10", s == 200 and sum(1 for a in acts if a["meetingTypeId"] == coproj["id"]) >= 7, str(len(acts) if s == 200 else acts))
s, decs = req("GET", f"/api/accounts/{ACC}/e/decision")
check("décisions reprises des sujets", s == 200 and len(decs) >= 4, str(len(decs) if s == 200 else decs))

# ---------------------------------------------------------------- relevé des actions
s, a = req("POST", f"/api/accounts/{ACC}/e/action", {"title": "Action test V15", "party": "CLIENT", "streamId": streams[0]["id"], "meetingTypeId": coproj["id"], "dueDate": "2026-10-20"})
check("création d'une action", s == 200 and a["status"] == "OPEN" and a["closedAt"] is None, str((s, a)))
s, a2 = req("PATCH", f"/api/accounts/{ACC}/e/action/{a['id']}", {"status": "DONE"})
check("action faite : date de clôture", s == 200 and a2["closedAt"], str(a2))
s, a3 = req("PATCH", f"/api/accounts/{ACC}/e/action/{a['id']}", {"status": "OPEN"})
check("action rouverte : date de clôture effacée", s == 200 and a3["closedAt"] is None)
s, r = req("POST", f"/api/accounts/{ACC}/e/action", {"title": "x", "party": "AUTRE"})
check("porteur invalide refusé", s == 400, str(s))
s, r = req("POST", f"/api/accounts/{ACC}/e/action", {"title": "x", "party": "WIFIRST", "streamId": "00000000-0000-0000-0000-000000000000"})
check("stream inconnu refusé", s in (400, 404), str(s))
s, r = req("POST", f"/api/accounts/{ACC}/e/action", {"title": "x", "party": "WIFIRST", "dueDate": "2026-02-30"})
check("échéance inexistante refusée", s == 400, str(s))
s, a4 = req("POST", f"/api/accounts/{ACC}/e/action", {"title": "Action close à la création", "party": "JOINT", "status": "CANCELLED"})
check("action créée abandonnée : date de clôture", s == 200 and a4["closedAt"], str(a4))
req("DELETE", f"/api/accounts/{ACC}/e/action/{a4['id']}")

# ---------------------------------------------------------------- registre des décisions
s, d = req("POST", f"/api/accounts/{ACC}/e/decision", {"title": "Décision test V15", "status": "PENDING", "streamId": streams[0]["id"], "meetingTypeId": coproj["id"]})
check("décision attendue créée", s == 200 and d["status"] == "PENDING", str((s, d)))
s, d2 = req("PATCH", f"/api/accounts/{ACC}/e/decision/{d['id']}", {"status": "TAKEN", "decidedOn": "2026-10-04"})
check("décision prise", s == 200 and d2["status"] == "TAKEN" and d2["decidedOn"] == "2026-10-04")
s, r = req("POST", f"/api/accounts/{ACC}/e/decision", {"title": "x", "status": "PEUT-ETRE"})
check("statut de décision invalide refusé", s == 400)
req("PATCH", f"/api/accounts/{ACC}/e/decision/{d['id']}", {"status": "PENDING", "decidedOn": None})

# ---------------------------------------------------------------- quoi de neuf depuis la séance précédente
s, ms = req("GET", f"/api/accounts/{ACC}/meetings?typeId={coproj['id']}")
last = ms[0]
s, ch = req("GET", f"/api/accounts/{ACC}/meetings/{last['id']}/changes")
check("quoi de neuf : réponse", s == 200 and {"since", "until", "cards", "streams", "actions", "decisions"} <= set(ch.keys()), str(s))
check("quoi de neuf : période cohérente", s == 200 and ch["since"] <= ch["until"] == last["date"], str((ch.get("since"), ch.get("until"))))
s, r = req("GET", f"/api/accounts/{ACC}/meetings/not-a-uuid/changes")
check("quoi de neuf : séance inconnue", s == 404, str(s))

# ---------------------------------------------------------------- revue de stream
s, rv = req("GET", f"/api/accounts/{ACC}/streams/{streams[0]['id']}/review")
check("revue de stream : réponse", s == 200 and rv["stream"]["id"] == streams[0]["id"], str(s))
check("revue de stream : action ouverte du stream", s == 200 and any(x["id"] == a["id"] for x in rv["actions"]))
check("revue de stream : décision attendue du stream", s == 200 and any(x["id"] == d["id"] for x in rv["decisions"]))
s, r = req("GET", f"/api/accounts/{ACC}/streams/00000000-0000-0000-0000-000000000000/review")
check("revue de stream : stream inconnu", s == 404, str(s))

# ---------------------------------------------------------------- bilan de sprint
s, sr = req("GET", f"/api/accounts/{ACC}/sprints/{sprint['id']}/review")
check("bilan de sprint : réponse", s == 200 and sr["sprint"]["id"] == sprint["id"] and sr["stats"]["total"] >= sr["stats"]["done"], str(s))
check("bilan de sprint (en cours) : restant = non terminés", s == 200 and sr["stats"]["carried"] == sr["stats"]["total"] - sr["stats"]["done"], str(sr.get("stats")))
done_sprint = next((x for x in boot["sprints"] if x["state"] == "DONE"), None)
if done_sprint:
    s, sr2 = req("GET", f"/api/accounts/{ACC}/sprints/{done_sprint['id']}/review")
    check("bilan d'un sprint terminé", s == 200, str(s))

# ---------------------------------------------------------------- import d'un CR d'atelier (Claude simulé)
reseau = next((x for x in streams if x["active"] and "seau" in x["name"].lower()), None)
s, ex = req("POST", f"/api/accounts/{ACC}/ai/extract", {"text": "Atelier IPAM. La Poste envoie le plan d'adressage avant le 15 octobre. Wifirst organise la recette. Le lot 1 démarre le 20.", "title": "Atelier IPAM", "date": "2026-10-03"})
check("extraction : réponse", s == 200 and len(ex["actions"]) == 2 and len(ex["cards"]) == 1 and len(ex["decisions"]) == 1, str((s, ex if s != 200 else "")))
if s == 200:
    check("extraction : porteur inconnu non rattaché", ex["actions"][0]["ownerId"] is None and ex["actions"][0]["ownerName"] == "Inconnu Test")
    check("extraction : échéance conservée", ex["actions"][0]["dueDate"] == "2026-10-15")
    check("extraction : stream inconnu ignoré", ex["cards"][0]["streamId"] is None)
    if reseau:
        check("extraction : stream reconnu", ex["actions"][0]["streamId"] == reseau["id"], str(ex["actions"][0]["streamId"]))
s, r = req("POST", f"/api/accounts/{ACC}/ai/extract", {"text": "trop court"})
check("extraction : texte trop court refusé", s == 400, str(s))
s, r = req("POST", f"/api/accounts/{ACC}/ai/extract", {"file": {"name": "cr.pdf", "base64": base64.b64encode(b"x" * 100).decode()}})
check("extraction : format non pris en charge", s == 400, str(s))
vtt = "WEBVTT\n\n1\n00:00:01.000 --> 00:00:04.000\nFlorent : on valide le démarrage du lot 1 le 20 octobre.\n\n2\n00:00:05.000 --> 00:00:09.000\nTuyen : La Poste envoie le plan IPAM avant le 15.\n"
s, ex2 = req("POST", f"/api/accounts/{ACC}/ai/extract", {"file": {"name": "visio.vtt", "base64": base64.b64encode(vtt.encode()).decode()}})
check("extraction : sous-titres .vtt", s == 200 and ex2["chars"] < len(vtt), str((s, ex2.get("chars") if isinstance(ex2, dict) else ex2)))
import docx

doc = docx.Document()
doc.add_paragraph("Compte rendu de l'atelier Réseau du 3 octobre 2026.")
doc.add_paragraph("Actions : La Poste envoie le plan IPAM avant le 15 octobre. Wifirst planifie la recette.")
bio = io.BytesIO()
doc.save(bio)
s, ex3 = req("POST", f"/api/accounts/{ACC}/ai/extract", {"file": {"name": "atelier.docx", "base64": base64.b64encode(bio.getvalue()).decode()}})
check("extraction : fichier Word", s == 200 and ex3["chars"] > 60, str((s, ex3 if s != 200 else ex3["chars"])))

# ---------------------------------------------------------------- brouillon de faits marquants
wm = req("GET", f"/api/accounts/{ACC}/meetings?typeId={weekly['id']}")[1]
if wm:
    s, sh = req("POST", f"/api/accounts/{ACC}/meetings/{wm[0]['id']}/suggest-highlights", {})
    check("faits marquants proposés", (s == 200 and len(sh["highlights"]) == 2) or (s == 400 and "Rien n'a changé" in str(sh)), str((s, sh)))

# ---------------------------------------------------------------- droits
viewer = client()
members = req("GET", f"/api/accounts/{ACC}/members")[1]
if not any(m.get("user", {}).get("email") == VIEWER[0] or m.get("email") == VIEWER[0] for m in members):
    req("POST", f"/api/accounts/{ACC}/members", {"email": VIEWER[0], "name": "Lecteur Tour", "password": VIEWER[1], "role": "VIEWER"})
login(viewer, *VIEWER)
s, r = viewer("POST", f"/api/accounts/{ACC}/e/action", {"title": "x", "party": "WIFIRST"})
check("lecteur : création d'action refusée", s == 403, str(s))
s, r = viewer("POST", f"/api/accounts/{ACC}/ai/extract", {"text": "x" * 100})
check("lecteur : extraction refusée", s == 403, str(s))
s, r = viewer("GET", f"/api/accounts/{ACC}/streams/{streams[0]['id']}/review")
check("lecteur : revue de stream lisible", s == 200, str(s))
s, tok = req("POST", "/api/auth/tokens", {"name": "test v15", "readOnly": True})
s, r = client()("POST", "/api/mcp", {"jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": "list_accounts", "arguments": {}}}, {"Authorization": f"Bearer {tok['token']}"})
check("jeton en lecture seule d'un super-administrateur : comptes listés", s == 200 and ACC in json.dumps(r) and "VIEWER" in json.dumps(r), str(r)[:200])
s, r = client()("GET", "/api/accounts", None, {"Authorization": f"Bearer {tok['token']}"})
check("jeton en lecture seule : liste REST des comptes", s == 200 and any(a["slug"] == ACC for a in r), str(r)[:200])
s, r = client()("POST", f"/api/accounts/{ACC}/ai/extract", {"text": "x" * 100}, {"Authorization": f"Bearer {tok['token']}"})
check("jeton d'accès : extraction refusée", s == 403, str(s))
req("DELETE", f"/api/auth/tokens/{tok['id']}")

# ---------------------------------------------------------------- slides des decks
from pptx import Presentation

all_sections = "cover,sprintReview,livrables,kanban,highlights,meteo,focus,expectations,coproj,statuses,topics,alerts,decisions,actions,planning"
s, deck = req("GET", f"/api/accounts/{ACC}/export/deck.pptx?sections={all_sections}&meetings={last['id']}&name=coproj")
try:
    prs = Presentation(io.BytesIO(deck))
    txt = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    tbl = " ".join(c.text for sl in prs.slides for sh in sl.shapes if sh.has_table for row in sh.table.rows for c in row.cells)
    check("deck complet lisible", len(prs.slides) > 8, f"{len(prs.slides)} slides")
    for title in ["Bilan du", "Les livrables du", "Actions en cours des streams", "Focus stream", "ce que nous attendons de", "Avancement des streams Wifirst, alertes et prérequis", "Registre des décisions", "Relevé des actions"]:
        check(f"slide « {title} »", title in txt, "")
    check("deck sans balisage brut", "**" not in txt + tbl)
    seeded = next(x["title"] for x in acts if x["meetingTypeId"] == coproj["id"] and x["status"] == "OPEN")
    check("deck : actions du COPROJ dans le relevé", seeded[:40] in tbl, seeded)
    check("deck : action postérieure à la séance absente", "Action test V15" not in tbl)
    overflow = [sh.name for sl in prs.slides for sh in sl.shapes if sh.left is not None and (sh.left + sh.width > prs.slide_width + 10 or sh.top + sh.height > prs.slide_height + 10)]
    check("deck : formes dans la slide", not overflow, str(overflow[:5]))
except Exception as e:
    check("deck complet lisible", False, repr(e)[:300])
s, deck2 = req("GET", f"/api/accounts/{ACC}/export/deck.pptx?sections=meetings")
try:
    check("section « meetings » toujours prise en charge", len(Presentation(io.BytesIO(deck2)).slides) >= 1)
except Exception as e:
    check("section « meetings » toujours prise en charge", False, repr(e))
s, deck3 = req("GET", f"/api/accounts/{ACC}/export/deck.pptx?sections=focus&focus={streams[0]['id']}")
try:
    p3 = Presentation(io.BytesIO(deck3))
    check("focus limité aux streams choisis", len(p3.slides) == 1, str(len(p3.slides)))
except Exception as e:
    check("focus limité aux streams choisis", False, repr(e))

# ---------------------------------------------------------------- export puis import : actions et décisions conservées
s, raw = req("GET", f"/api/accounts/{ACC}/export/account.json")
data = raw if isinstance(raw, dict) else json.loads(raw)
check("export : actions et décisions", len(data.get("actions", [])) == len(req("GET", f"/api/accounts/{ACC}/e/action")[1]) and len(data.get("decisions", [])) >= 4)
c = copy.deepcopy(data)
c["account"]["slug"] = "test-reimport-v15"
c["account"]["name"] = "Test réimport V15"
s, r = req("POST", "/api/admin/import", c)
check("réimport avec actions et décisions", s == 200, str((s, r if s != 200 else "")))
if s == 200:
    d2 = req("GET", "/api/accounts/test-reimport-v15/export/account.json")[1]
    d2 = d2 if isinstance(d2, dict) else json.loads(d2)
    check("réimport : mêmes actions", len(d2["actions"]) == len(data["actions"]), f"{len(d2['actions'])} / {len(data['actions'])}")
    check("réimport : mêmes décisions", len(d2["decisions"]) == len(data["decisions"]))
    linked = [x for x in data["decisions"] if x.get("topic") is not None]
    linked2 = [x for x in d2["decisions"] if x.get("topic") is not None]
    check("réimport : décisions liées aux sujets", len(linked) == len(linked2), f"{len(linked2)} / {len(linked)}")
    subprocess.run(["psql", DB, "-qc", "delete from accounts where slug='test-reimport-v15'"], check=True)

# nettoyage
req("DELETE", f"/api/accounts/{ACC}/e/action/{a['id']}")
req("DELETE", f"/api/accounts/{ACC}/e/decision/{d['id']}")
print("RESULT", "OK" if not fails else f"{len(fails)} FAIL: {fails}")
