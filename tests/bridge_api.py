"""WiBridge : tests de l'API locale (comptes et séparation avec WacMan, droits par stream, échanges, fichiers, notifications, export)."""
import os, io, json, time, uuid, urllib.request, urllib.error, urllib.parse, http.cookiejar

ENV_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "apps", "api", ".env.dev")
B = "http://localhost:4000"
PWD = open(ENV_FILE).read().split("BOOTSTRAP_ADMIN_PASSWORD=")[1].split("\n")[0].strip()
ADMIN = "florent@omgt.fr"
RUN = uuid.uuid4().hex[:6]  # adresses propres à chaque exécution
UPWD = "Bridge-Test-2026"
fails = []


def check(name, cond, extra=""):
    print(("OK  " if cond else "FAIL"), name, extra if not cond else "")
    if not cond:
        fails.append(name)


class S:
    """Session HTTP (cookies propres)."""

    def __init__(self):
        self.cj = http.cookiejar.CookieJar()
        self.op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.cj))

    def req(self, method, path, body=None, headers=None, raw=None, base=B):
        data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
        h = {**({"Content-Type": "application/json"} if body is not None else {}), **(headers or {})}
        r = urllib.request.Request(base + path, data=data, method=method, headers=h)
        try:
            with self.op.open(r) as res:
                t = res.read()
                ct = res.headers.get("content-type", "")
                self.last_headers = res.headers
                return res.status, (json.loads(t) if "json" in ct and t else t)
        except urllib.error.HTTPError as e:
            t = e.read().decode()
            self.last_headers = e.headers
            try:
                return e.code, json.loads(t)
            except Exception:
                return e.code, t

    def cookie(self, name):
        return next((c.value for c in self.cj if c.name == name), None)



def blogin(s, email, pwd=UPWD, trust=True):
    st, r = s.req("POST", "/api/bridge/auth/login", {"email": email, "password": pwd})
    assert st == 200, (email, st, r)
    if r.get("trusted"):
        return "trusted"
    st, r = s.req("POST", "/api/bridge/auth/verify", {"challengeId": r["challengeId"], "code": r["devCode"], "trust": trust})
    assert st == 200, (email, st, r)
    return "code"


def wlogin(s, email, pwd):
    st, r = s.req("POST", "/api/auth/login", {"email": email, "password": pwd})
    if st != 200:
        return st
    st, r = s.req("POST", "/api/auth/verify", {"challengeId": r["challengeId"], "code": r["devCode"]})
    return st


def outbox(after=0):
    """E-mails de la boîte d'envoi de développement, après le repère donné (numéro du dernier e-mail lu)."""
    return S().req("GET", f"/api/bridge/dev/outbox?after={after}")[1]


def mark():
    o = outbox()
    return o[-1]["n"] if o else 0


# ---------------------------------------------------------------------------
# Super-administrateur, client et streams
# ---------------------------------------------------------------------------
su = S()
check("connexion super-admin (code sur nouvel appareil)", blogin(su, ADMIN, PWD) == "code")
check("seconde connexion : appareil de confiance, pas de code", blogin(su, ADMIN, PWD) == "trusted")
st, boot = su.req("GET", "/api/bridge/c/la-poste")
check("client La Poste", st == 200 and boot["client"]["clientName"] == "La Poste", str(boot)[:200])
names = [s["name"] for s in boot["streams"]]
check("7 streams par défaut dans l'ordre", names[:7] == ["Technico-fonctionnelle", "Sécurité", "Déploiement", "Mainteneur Postal", "Exploitation", "Outillage & Data", "Gouvernance"], str(names))
SID = {s["name"]: s["id"] for s in boot["streams"]}
CID = boot["client"]["id"]
check("réglages par défaut", boot["settings"] == {"providerAnswersForClient": True, "providerEditsAll": True, "clientCloseScope": "ANY", "clientCanReopen": True, "notifications": True}, str(boot["settings"]))
base = "/api/bridge/c/la-poste"


def invite(email, name, side, default, stream_access=None, wacman=False):
    st, r = su.req("POST", "/api/bridge/admin/users", {"email": email, "name": name, "wacmanAccess": wacman, "memberships": [{"clientId": CID, "side": side, "defaultAccess": default, "streamAccess": stream_access or {}}]})
    assert st == 200, (st, r)
    return r


def accept(r, name, pwd=UPWD):
    tok = r["devLink"].split("#")[1]
    s = S()
    st, info = s.req("POST", "/api/bridge/auth/invitation/info", {"token": tok})
    st2, res = s.req("POST", "/api/bridge/auth/invitation/accept", {"token": tok, "name": name, "password": pwd})
    return s, st, info, st2, res, tok


WF, LP, RD, SEC, BOTH = (f"wf.{RUN}@wibridge.test", f"lp.{RUN}@wibridge.test", f"rd.{RUN}@wibridge.test", f"sec.{RUN}@wibridge.test", f"both.{RUN}@wibridge.test")
r = invite(WF, "Wendy Wifirst", "PROVIDER", "PROVIDER")
check("invitation : mode invite et lien", r["mode"] == "invite" and r["sent"] and "/invitation#" in r.get("devLink", ""), str(r))
mails = [m for m in outbox() if m["to"] == WF]
check("e-mail d'invitation envoyé", bool(mails) and "Invitation" in mails[-1]["subject"] and "/invitation#" in mails[-1]["text"], str(mails[-1:])[:300])
check("e-mail d'invitation : notifications à activer dans Mon compte", bool(mails) and "Mon compte" in mails[-1]["text"] and "ne vous envoie pas d'e-mail" in mails[-1]["text"])
tok = r["devLink"].split("#")[1]
x = S()
st, res = x.req("POST", "/api/bridge/auth/invitation/accept", {"token": tok, "name": "Wendy", "password": "court"})
check("invitation : mot de passe trop faible refusé", st == 400, str(res))
wf, st, info, st2, res, _ = accept(r, "Wendy Wifirst")
check("invitation : informations", st == 200 and info["email"] == WF and info["hasPassword"] is False and info["clients"] == ["La Poste"], str(info))
check("invitation acceptée, session et appareil de confiance", st2 == 200 and wf.cookie("wib_session") and wf.cookie("wib_device"), str(res))
st, res = S().req("POST", "/api/bridge/auth/invitation/accept", {"token": tok, "name": "X", "password": UPWD})
check("invitation : lien à usage unique", st == 400, str(res))

lp, *_ = accept(invite(LP, "Louise LaPoste", "CLIENT", "CLIENT"), "Louise LaPoste")
rd, *_ = accept(invite(RD, "Rémi Lecteur", "CLIENT", "READ"), "Rémi Lecteur")
sec, *_ = accept(invite(SEC, "Sacha Sécurité", "CLIENT", "NONE", {SID["Sécurité"]: "CLIENT"}), "Sacha Sécurité")
both, *_ = accept(invite(BOTH, "Bastien Deux", "PROVIDER", "BOTH"), "Bastien Deux")
st, me = lp.req("GET", "/api/bridge/auth/me")
check("préférence par défaut : aucun e-mail", st == 200 and next(c for c in me["clients"] if c["slug"] == "la-poste")["notify"] == "NONE", str(me)[:300])

# connexion sur un nouvel appareil : code, puis appareil de confiance ; sans confiance, le code revient
d1 = S()
check("nouvel appareil : code demandé", blogin(d1, LP) == "code")
check("même appareil : plus de code", blogin(d1, LP) == "trusted")
d2 = S()
blogin(d2, LP, trust=False)
check("appareil non retenu : pas de cookie d'appareil", d2.cookie("wib_device") is None)
check("appareil non retenu : code redemandé", blogin(d2, LP, trust=False) == "code")
st, r = S().req("POST", "/api/bridge/auth/login", {"email": LP, "password": "mauvais-mot-de-passe-1"})
check("mot de passe faux : 401 générique", st == 401 and "incorrect" in r["error"], str(r))
st, r = S().req("POST", "/api/bridge/auth/login", {"email": f"inconnu.{RUN}@wibridge.test", "password": UPWD})
check("compte inconnu : même message", st == 401 and "incorrect" in r["error"], str(r))
x = S()
st, r = x.req("POST", "/api/bridge/auth/login", {"email": LP, "password": UPWD})
st, r2 = x.req("POST", "/api/bridge/auth/verify", {"challengeId": r["challengeId"], "code": "000000" if r["devCode"] != "000000" else "111111"})
check("code incorrect", st == 401 and "incorrect" in r2["error"], str(r2))

# ---------------------------------------------------------------------------
# Séparation des comptes WiBridge et WacMan
# ---------------------------------------------------------------------------
check("compte WiBridge : connexion WacMan refusée", wlogin(S(), LP, UPWD) == 401)
st, r = S().req("POST", "/api/auth/forgot", {"email": LP})
check("compte WiBridge : pas de code de réinitialisation WacMan", st == 200 and "devCode" not in r, str(r))
st, r = lp.req("GET", "/api/auth/me")
check("session WiBridge sans effet sur WacMan", st == 401, str(r))
st, r = lp.req("GET", "/api/accounts")
check("session WiBridge : comptes WacMan refusés", st == 401, str(r))
w = S()
check("super-admin : connexion WacMan", wlogin(w, ADMIN, PWD) == 200)
st, r = w.req("GET", "/api/bridge/auth/me")
check("session WacMan sans effet sur WiBridge", st == 401, str(r))
# jetons forgés : le jeton d'une application présenté à l'autre est refusé
check("jeton WacMan présenté comme session WiBridge : refusé", S().req("GET", "/api/bridge/auth/me", headers={"Cookie": f"wib_session={w.cookie('wacman_session')}"})[0] == 401)
check("jeton WiBridge présenté comme session WacMan : refusé", S().req("GET", "/api/auth/me", headers={"Cookie": f"wacman_session={su.cookie('wib_session')}"})[0] == 401)
check("contrôle : vrais cookies acceptés", S().req("GET", "/api/bridge/auth/me", headers={"Cookie": f"wib_session={su.cookie('wib_session')}"})[0] == 200)
st, users = w.req("GET", "/api/admin/users")
check("administration WacMan : comptes WiBridge absents", st == 200 and not any(u["email"] == LP for u in users), str([u["email"] for u in users])[:200])
# le super-administrateur ouvre WacMan à un compte WiBridge
st, users = su.req("GET", "/api/bridge/admin/users")
wf_user = next(u for u in users if u["email"] == WF)
check("liste WiBridge : accès WiBridge seul", wf_user["bridgeAccess"] and not wf_user["wacmanAccess"] and wf_user["passwordSet"], str(wf_user)[:300])
st, r = su.req("PATCH", f"/api/bridge/admin/users/{wf_user['id']}", {"wacmanAccess": True})
check("accès WacMan ouvert depuis WiBridge", st == 200, str(r))
check("compte WiBridge avec accès WacMan : connexion WacMan", wlogin(S(), WF, UPWD) == 200)
st, users = w.req("GET", "/api/admin/users")
check("administration WacMan : le compte apparaît", any(u["email"] == WF for u in users))
# compte WacMan existant invité dans WiBridge : simple ouverture d'accès, mot de passe conservé
WAC = f"wacman.{RUN}@wibridge.test"
st, r = w.req("POST", "/api/admin/users", {"email": WAC, "name": "Wanda WacMan", "password": "Wacman-Test-2026"})
check("compte WacMan créé", st == 200, str(r))
check("compte WacMan : pas d'accès WiBridge", S().req("POST", "/api/bridge/auth/login", {"email": WAC, "password": "Wacman-Test-2026"})[0] == 401)
# mot de passe oublié sans compte WiBridge : pas de code, mais un e-mail qui explique comment demander un accès
for addr, who in ((WAC, "compte WacMan seul"), (f"inconnu.{RUN}@wibridge.test", "adresse inconnue")):
    before = mark()
    st, r = S().req("POST", "/api/bridge/auth/forgot", {"email": addr.upper()})
    got = [m for m in outbox(before) if m["to"] == addr]
    check(f"mot de passe oublié, {who} : réponse identique, sans code", st == 200 and r.get("challengeId") and "devCode" not in r, str(r))
    check(f"mot de passe oublié, {who} : e-mail d'explication", len(got) == 1 and "demande de nouveau mot de passe" in got[0]["subject"] and "invitation" in got[0]["text"], str(got)[:300])
r = invite(WAC, "Wanda WacMan", "PROVIDER", "READ")
check("invitation d'un compte WacMan : mode accès", r["mode"] == "access" and r["devLink"].endswith("/login"), str(r))
check("e-mail d'accès envoyé", any(m["to"] == WAC and "accès" in m["subject"].lower() for m in outbox()))
check("compte WacMan : connexion WiBridge avec son mot de passe", blogin(S(), WAC, "Wacman-Test-2026") == "code")
# WacMan : un administrateur de compte ne peut pas ouvrir WacMan à un compte WiBridge ; le super-administrateur oui
st, accs = w.req("GET", "/api/accounts")
if not accs:
    w.req("POST", "/api/accounts", {"name": f"Compte test {RUN}", "clientName": "Test", "clientShortName": "TST"})
    st, accs = w.req("GET", "/api/accounts")
if accs:
    acc = accs[0]["slug"]
    st, r = w.req("POST", f"/api/accounts/{acc}/members", {"email": RD, "name": "Rémi Lecteur", "role": "VIEWER"})
    check("WacMan : super-admin ouvre WacMan à un compte WiBridge", st == 200, str(r))
    check("compte ouvert : connexion WacMan", wlogin(S(), RD, UPWD) == 200)
    su.req("PATCH", f"/api/bridge/admin/users/{next(u for u in su.req('GET', '/api/bridge/admin/users')[1] if u['email'] == RD)['id']}", {"wacmanAccess": False})
else:
    print("    (aucun compte WacMan local : test d'ajout de membre ignoré)")

# ---------------------------------------------------------------------------
# Questions et droits
# ---------------------------------------------------------------------------
def create(s, subject, streams, **kw):
    return s.req("POST", f"{base}/questions", {"subject": subject, "body": kw.pop("body", ""), "streamIds": [SID[x] for x in streams], **kw})


st, q1 = create(wf, f"Liste des ATM {RUN}", ["Déploiement"], body="Pouvez-vous nous fournir la liste des ATM sous la forme d'un fichier **Excel** ?", assignedParty="CLIENT", dueDate="2026-10-30")
check("Wifirst pose une question au client", st == 200 and q1["assignedParty"] == "CLIENT" and q1["askedByParty"] == "PROVIDER" and q1["status"] == "OPEN" and q1["askedBy"]["name"] == "Wendy Wifirst", str(q1)[:300])
check("numérotation", isinstance(q1.get("ref"), int))
st, r = create(lp, f"Calendrier {RUN}", ["Déploiement"], assignedParty="CLIENT")
check("le client ne s'attribue pas sa propre question", st == 400, str(r))
st, q2 = create(lp, f"Calendrier des interventions {RUN}", ["Déploiement", "Exploitation"], body="Merci de nous transmettre le calendrier.")
check("question du client attribuée à Wifirst", st == 200 and q2["assignedParty"] == "PROVIDER" and q2["askedByParty"] == "CLIENT", str(q2)[:300])
check("question multi-streams", st == 200 and len(q2["streamIds"]) == 2)
st, r = create(rd, "Lecteur", ["Déploiement"])
check("lecteur : création refusée", st == 403, str(r))
st, r = create(sec, "Hors droits", ["Déploiement"])
check("droit limité à Sécurité : création ailleurs refusée", st == 403, str(r))
st, q3 = create(sec, f"Tests d'intrusion {RUN}", ["Sécurité"])
check("droit limité à Sécurité : création dans Sécurité", st == 200, str(q3)[:200])
st, r = create(wf, " ", ["Déploiement"])
check("sujet obligatoire", st == 400, str(r))
st, r = create(wf, "Date", ["Déploiement"], dueDate="2026-02-30")
check("date inexistante refusée", st == 400, str(r))
st, r = wf.req("POST", f"{base}/questions", {"subject": "x", "streamIds": [str(uuid.uuid4())]})
check("stream inconnu refusé", st == 400, str(r))

# visibilité par stream
st, lst = sec.req("GET", f"{base}/questions")
refs = {q["ref"] for q in lst}
check("visibilité : seulement les streams ouverts", q3["ref"] in refs and q1["ref"] not in refs and q2["ref"] not in refs, str(refs))
st, r = sec.req("GET", f"{base}/questions/{q1['id']}")
check("question hors droits : introuvable", st == 404, str(r))
st, r = sec.req("GET", f"{base}/questions/{q1['ref']}")
check("question hors droits par numéro : introuvable", st == 404)
st, rq = rd.req("GET", f"{base}/questions/{q1['ref']}")
check("lecteur : lecture par numéro", st == 200 and rq["id"] == q1["id"])
p = rq["perms"]
check("lecteur : aucun droit d'écriture", not p["edit"] and not p["respondAs"] and not p["close"] and not p["reassign"] and not p["delete"] and not p["streams"], str(p))

# échanges : le client répond à une question qui lui est attribuée, pas à celle attribuée à Wifirst
st, r = lp.req("POST", f"{base}/questions/{q2['id']}/messages", {"body": "Précision", "outcome": "PROVIDER"})
check("client : réponse refusée sur une question attribuée à Wifirst", st == 403, str(r))
st, r = lp.req("POST", f"{base}/questions/{q1['id']}/messages", {"body": "Premiers éléments, fichier complet la semaine prochaine.", "outcome": "CLIENT"})
check("client : réponse partielle, attribution conservée", st == 200 and r["status"] == "IN_PROGRESS" and r["assignedParty"] == "CLIENT", str(r)[:200])
m1 = r["messages"][-1]
check("message : organisation et issue", m1["party"] == "CLIENT" and m1["outcome"] == "ASSIGN" and m1["assignedBefore"] == "CLIENT" and m1["assignedAfter"] == "CLIENT")
check("message : modifiable et supprimable par son auteur", m1["perms"]["edit"] and m1["perms"]["delete"])
st, r = lp.req("PATCH", f"{base}/messages/{m1['id']}", {"body": "Premiers éléments (extraction partielle), fichier complet la semaine prochaine."})
check("client : modifie son dernier message", st == 200 and r["messages"][-1]["editedAt"], str(r)[:200])
st, r = wf.req("POST", f"{base}/questions/{q1['id']}/messages", {"body": "Merci, nous attendons le fichier complet.", "outcome": "CLIENT"})
check("Wifirst : relance, attribution et statut inchangés", st == 200 and r["assignedParty"] == "CLIENT" and r["status"] == "IN_PROGRESS", str(r)[:200])
st, r = lp.req("PATCH", f"{base}/messages/{m1['id']}", {"body": "Premiers éléments (extraction partielle), fichier complet sous huit jours."})
check("client : modifie son message même après une réponse", st == 200 and r["messages"][-2]["body"].endswith("sous huit jours."), str(r)[:200])
st, r = wf.req("PATCH", f"{base}/messages/{m1['id']}", {"body": "Premiers éléments reçus par e-mail le 07/10."})
check("Wifirst : modifie un message du client (peut tout changer)", st == 200, str(r)[:200])
st, r = lp.req("POST", f"{base}/questions/{q1['id']}/messages", {"body": "Voici le fichier complet.", "outcome": "PROVIDER"})
check("client : réponse et attribution à Wifirst", st == 200 and r["assignedParty"] == "PROVIDER" and r["status"] == "OPEN", str(r)[:200])
st, r = lp.req("POST", f"{base}/questions/{q1['id']}/messages", {"body": "Encore un mot", "outcome": "CLIENT"})
check("client : plus de réponse une fois attribuée à Wifirst", st == 403, str(r))
st, r = wf.req("POST", f"{base}/questions/{q1['id']}/messages", {"body": "Réponse suffisante, merci.", "outcome": "CLOSE"})
check("Wifirst : réponse et clôture", st == 200 and r["status"] == "CLOSED" and r["closedByName"] == "Wendy Wifirst" and r["messages"][-1]["outcome"] == "CLOSE", str(r)[:200])
st, r = wf.req("POST", f"{base}/questions/{q1['id']}/messages", {"body": "x", "outcome": "CLIENT"})
check("question clôturée : réponse refusée", st == 400, str(r))
st, r = rd.req("POST", f"{base}/questions/{q1['id']}/reopen", {})
check("lecteur : réouverture refusée", st == 403, str(r))
st, r = lp.req("POST", f"{base}/questions/{q1['id']}/reopen", {"body": "Il manque les ATM de Corse."})
check("client : rouvre, attribuée à Wifirst par défaut", st == 200 and r["status"] == "OPEN" and r["assignedParty"] == "PROVIDER" and r["messages"][-1]["outcome"] == "REOPEN", str(r)[:200])

# règles de clôture et de réouverture du client
st, r = lp.req("POST", f"{base}/questions/{q1['id']}/close", {})
check("client : clôture d'une question attribuée à Wifirst (toutes les questions)", st == 200 and r["status"] == "CLOSED", str(r)[:200])
su.req("PATCH", f"/api/bridge/admin/clients/{CID}", {"settings": {"clientCanReopen": False, "clientCloseScope": "OWN_OR_ASSIGNED"}})
st, r = lp.req("POST", f"{base}/questions/{q1['id']}/reopen", {})
check("réglage : réouverture par le client désactivée", st == 403, str(r))
st, r = wf.req("POST", f"{base}/questions/{q1['id']}/reopen", {"party": "PROVIDER"})
check("Wifirst : rouvre et s'attribue la question", st == 200 and r["assignedParty"] == "PROVIDER", str(r)[:200])
st, r = lp.req("POST", f"{base}/questions/{q1['id']}/close", {})
check("réglage : clôture limitée aux questions du client", st == 403, str(r))
st, r = lp.req("POST", f"{base}/questions/{q2['id']}/close", {})
check("réglage : le client clôt sa propre question", st == 200, str(r)[:200])
su.req("PATCH", f"/api/bridge/admin/clients/{CID}", {"settings": {"clientCanReopen": True, "clientCloseScope": "ANY"}})
lp.req("POST", f"{base}/questions/{q2['id']}/reopen", {})

# Wifirst répond à la place du client, sauf réglage contraire
st, q4 = create(wf, f"Contacts sécurité {RUN}", ["Sécurité"], assignedParty="CLIENT")
st, r = wf.req("POST", f"{base}/questions/{q4['id']}/messages", {"body": "Réponse reçue par e-mail : la DSEM.", "outcome": "PROVIDER"})
check("Wifirst répond à une question attribuée au client", st == 200 and r["messages"][-1]["party"] == "PROVIDER", str(r)[:200])
wf.req("POST", f"{base}/questions/{q4['id']}/assign", {"party": "CLIENT"})
su.req("PATCH", f"/api/bridge/admin/clients/{CID}", {"settings": {"providerAnswersForClient": False, "providerEditsAll": False}})
st, r = wf.req("POST", f"{base}/questions/{q4['id']}/messages", {"body": "x", "outcome": "CLIENT"})
check("réglage : Wifirst ne répond plus à la place du client", st == 403, str(r))
st, r = wf.req("PATCH", f"{base}/questions/{q2['id']}", {"subject": "Changé par Wifirst"})
check("réglage : Wifirst ne modifie plus les questions du client", st == 403, str(r))
su.req("PATCH", f"/api/bridge/admin/clients/{CID}", {"settings": {"providerAnswersForClient": True, "providerEditsAll": True}})
st, r = wf.req("PATCH", f"{base}/questions/{q2['id']}", {"subject": f"Calendrier des interventions sur les bureaux pilotes {RUN}"})
check("Wifirst modifie le sujet d'une question du client", st == 200 and r["subject"].startswith("Calendrier des interventions sur"), str(r)[:200])
st, r = lp.req("PATCH", f"{base}/questions/{q2['id']}", {"subject": f"Calendrier des interventions {RUN}"})
check("le créateur modifie le sujet", st == 200, str(r)[:200])
st, r = lp.req("PATCH", f"{base}/questions/{q1['id']}", {"subject": "Pas à moi"})
check("le client ne modifie pas le sujet d'une question de Wifirst", st == 403, str(r))
st, r = lp.req("PATCH", f"{base}/questions/{q2['id']}", {"streamIds": [SID["Déploiement"], SID["Exploitation"], SID["Gouvernance"]]})
check("créateur : ajoute un stream", st == 200 and len(r["streamIds"]) == 3, str(r)[:200])
st, r = lp.req("PATCH", f"{base}/questions/{q2['id']}", {"streamIds": []})
check("au moins un stream", st == 400, str(r))
st, r = sec.req("PATCH", f"{base}/questions/{q3['id']}", {"streamIds": [SID["Sécurité"], SID["Déploiement"]]})
check("créateur : pas d'ajout d'un stream hors de ses droits", st == 403, str(r))
st, r = lp.req("PATCH", f"{base}/questions/{q2['id']}", {"dueDate": "2026-11-15"})
check("échéance modifiée", st == 200 and r["dueDate"] == "2026-11-15")

# éditeur des deux : répond au nom du client
st, q5 = create(both, f"Question au nom du client {RUN}", ["Gouvernance"], askedByParty="CLIENT")
check("éditeur des deux : question au nom du client", st == 200 and q5["askedByParty"] == "CLIENT" and q5["assignedParty"] == "PROVIDER", str(q5)[:200])
st, r = both.req("POST", f"{base}/questions/{q4['id']}/messages", {"body": "Réponse de La Poste transmise en séance.", "party": "CLIENT", "outcome": "PROVIDER"})
check("éditeur des deux : répond au nom du client", st == 200 and r["messages"][-1]["party"] == "CLIENT", str(r)[:200])
st, r = lp.req("POST", f"{base}/questions/{q2['id']}/messages", {"body": "x", "party": "PROVIDER", "outcome": "PROVIDER"})
check("client : pas de réponse au nom de Wifirst", st == 403, str(r))

# changement d'attribution sans message
st, q6 = create(wf, f"Attribution {RUN}", ["Exploitation"], assignedParty="CLIENT")
st, r = lp.req("POST", f"{base}/questions/{q6['id']}/assign", {"party": "PROVIDER"})
check("client : renvoie à Wifirst une question qui lui est attribuée", st == 200 and r["assignedParty"] == "PROVIDER", str(r)[:200])
st, r = lp.req("POST", f"{base}/questions/{q6['id']}/assign", {"party": "CLIENT"})
check("client : ne reprend pas une question attribuée à Wifirst", st == 403, str(r))

# suppression d'une réponse par son auteur : texte retiré, issue et attribution conservées, trace à l'historique
st, q7 = create(wf, f"Suppression {RUN}", ["Sécurité"], assignedParty="CLIENT")
st, r = lp.req("POST", f"{base}/questions/{q7['id']}/messages", {"body": "Réponse à retirer", "outcome": "PROVIDER"})
mid = r["messages"][-1]["id"]
check("message : supprimable par son auteur", st == 200 and r["messages"][-1]["perms"]["delete"], str(r)[:200])
st, r = sec.req("DELETE", f"{base}/messages/{mid}")
check("autre membre du client : suppression refusée", st == 403, str(r))
st, r = rd.req("DELETE", f"{base}/messages/{mid}")
check("lecteur : suppression refusée", st in (403, 404), str(r))
st, r = lp.req("DELETE", f"{base}/messages/{mid}")
dm = r["messages"][-1] if st == 200 else {}
check("auteur : réponse supprimée, texte retiré, issue conservée", st == 200 and dm["deletedAt"] and dm["body"] == "" and dm["deletedByName"] == "Louise LaPoste" and dm["outcome"] == "ASSIGN" and not dm["perms"]["edit"], str(r)[:300])
check("suppression : attribution et statut inchangés, échange décompté", r["assignedParty"] == "PROVIDER" and r["status"] == "OPEN" and r["messageCount"] == 0 and r["lastMessage"] is None, str(r)[:300])
st, r = lp.req("PATCH", f"{base}/messages/{mid}", {"body": "Retour"})
check("message supprimé : plus modifiable", st == 404, str(r))
st, r = lp.req("DELETE", f"{base}/messages/{mid}")
check("message supprimé : pas de seconde suppression", st == 404, str(r))
st, h = lp.req("GET", f"{base}/questions/{q7['id']}/history")
ev = next((e for e in h if e["action"] == "message_delete"), None)
check("historique : suppression et texte supprimé", bool(ev) and ev["changes"].get("message") == "Réponse à retirer" and ev["userName"] == "Louise LaPoste", str(ev)[:300])
st, r = lp.req("GET", f"{base}/search?q=" + urllib.parse.quote("retirer"))
check("recherche : rien dans une réponse supprimée", st == 200 and q7["id"] not in r["ids"], str(r))
wf.req("POST", f"{base}/questions/{q7['id']}/assign", {"party": "CLIENT"})
st, r = lp.req("POST", f"{base}/questions/{q7['id']}/messages", {"body": "Nouvelle réponse", "outcome": "PROVIDER"})
st, r = wf.req("DELETE", f"{base}/messages/{r['messages'][-1]['id']}")
check("Wifirst : supprime une réponse du client (peut tout changer)", st == 200, str(r)[:200])

# historique
st, h = lp.req("GET", f"{base}/questions/{q1['id']}/history")
acts = [e["action"] for e in h]
check("historique : toutes les étapes dans l'ordre", st == 200 and acts[0] == "create" and "answer" in acts and "message_edit" in acts and "close" in acts and "reopen" in acts, str(acts))
ev = next(e for e in h if e["action"] == "message_edit")
check("historique : avant et après", isinstance(ev["changes"].get("message"), list) and len(ev["changes"]["message"]) == 2, str(ev)[:300])
st, h2 = lp.req("GET", f"{base}/questions/{q2['id']}/history")
up = [e for e in h2 if e["action"] == "update" and "subject" in e["changes"]]
check("historique : sujet modifié par Wifirst, auteur et valeurs", bool(up) and up[0]["userName"] == "Wendy Wifirst" and up[0]["changes"]["subject"][1].startswith("Calendrier des interventions sur"), str(up)[:300])
check("historique : streams modifiés en noms", any("streams" in e["changes"] and "Gouvernance" in e["changes"]["streams"][1] for e in h2))

# suppression et corbeille
st, q7 = create(lp, f"À supprimer {RUN}", ["Gouvernance"])
st, r = wf.req("DELETE", f"{base}/questions/{q7['id']}")
check("suppression : réservée au créateur ou au super-admin", st == 403, str(r))
st, r = lp.req("DELETE", f"{base}/questions/{q7['id']}")
check("suppression par le créateur (sans réponse)", st == 200)
check("question supprimée : retirée de la liste", all(q["id"] != q7["id"] for q in lp.req("GET", f"{base}/questions")[1]))
check("corbeille : visible du super-admin", any(q["id"] == q7["id"] for q in su.req("GET", f"{base}/questions?deleted=1")[1]))
st, r = su.req("POST", f"{base}/questions/{q7['id']}/restore", {})
check("restauration", st == 200 and r["deletedAt"] is None, str(r)[:200])
st, r = wf.req("DELETE", f"{base}/questions/{q1['id']}")
check("pas de suppression d'une question avec échanges, même par son créateur", st == 403, str(r))

# ---------------------------------------------------------------------------
# Pièces jointes
# ---------------------------------------------------------------------------
st, t = lp.req("GET", "/api/bridge/auth/upload-token")
check("jeton de dépôt", st == 200 and t["token"], str(t))
content = b"PK\x03\x04" + os.urandom(2048)


def upload(token, data, name, qs="", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", s=None):
    return (s or S()).req("POST", f"{base}/files{qs}", raw=data, headers={"Authorization": f"Bearer {token}", "Content-Type": "application/octet-stream", "X-File-Name": urllib.parse.quote(name), "X-File-Type": mime})


st, f = upload(t["token"], content, "Liste des ATM é.xlsx")
check("dépôt en attente", st == 200 and not f["attached"] and f["size"] == len(content) and f["name"] == "Liste des ATM é.xlsx", str(f))
st, r = lp.req("POST", f"{base}/questions/{q6['id']}/messages", {"body": "", "outcome": "PROVIDER", "fileIds": [f["id"]]})
check("réponse avec pièce jointe seule refusée (question attribuée à Wifirst)", st == 403)
wf.req("POST", f"{base}/questions/{q6['id']}/assign", {"party": "CLIENT"})
st, r = lp.req("POST", f"{base}/questions/{q6['id']}/messages", {"body": "Le fichier demandé.", "outcome": "PROVIDER", "fileIds": [f["id"]]})
check("réponse avec pièce jointe", st == 200 and len(r["files"]) == 1 and r["files"][0]["messageId"] == r["messages"][-1]["id"], str(r)[:300])
st, lst = lp.req("GET", f"{base}/questions")
check("liste : nombre de pièces jointes", next(q for q in lst if q["id"] == q6["id"])["fileCount"] == 1)
st, link = rd.req("GET", f"{base}/files/{f['id']}/link")
check("lien de téléchargement (lecteur)", st == 200 and link["path"].startswith("/api/bridge/dl/"), str(link))
x = S()
st, data = x.req("GET", link["path"])
check("téléchargement : contenu identique", st == 200 and data == content, f"{st} {len(data) if isinstance(data, bytes) else data}")
check("téléchargement : en pièce jointe, nom conservé", "attachment" in x.last_headers.get("content-disposition", "") and "Liste%20des%20ATM" in x.last_headers.get("content-disposition", ""), x.last_headers.get("content-disposition"))
st, r = S().req("GET", "/api/bridge/dl/abc.def.ghi")
check("lien invalide refusé", st == 403, str(r))
st, r = sec.req("GET", f"{base}/files/{f['id']}/link")
check("pièce jointe hors droits : introuvable", st == 404, str(r))
st, t2 = rd.req("GET", "/api/bridge/auth/upload-token")
st, r = upload(t2["token"], b"abc", "a.txt", mime="text/plain")
check("lecteur : dépôt refusé", st == 403, str(r))
st, r = upload(t["token"], b"MZ" + os.urandom(100), "outil.exe", mime="application/octet-stream")
check("exécutable refusé", st == 400, str(r))
st, r = upload(t["token"], os.urandom(21 * 1024 * 1024), "gros.bin", mime="application/octet-stream")
check("fichier de plus de 20 Mo refusé", st in (400, 413), str(r)[:120])
st, r = upload("pas-un-jeton", b"abc", "a.txt")
check("dépôt sans jeton valide refusé", st == 401, str(r)[:120])
st, r = lp.req("POST", f"{base}/files", raw=b"abc", headers={"Content-Type": "application/octet-stream", "X-File-Name": "a.txt"})
check("dépôt par cookie de session (relais) accepté", st == 200, str(r)[:120])
# ajout direct à la question par son créateur, retrait
st, f2 = upload(t["token"], b"%PDF-1.4 test", "note.pdf", qs=f"?questionId={q2['id']}", mime="application/pdf")
check("pièce jointe ajoutée à la question", st == 200 and f2["attached"], str(f2))
st, r = rd.req("GET", f"{base}/questions/{q2['id']}")
fx = next(x for x in r["files"] if x["id"] == f2["id"])
check("lecteur : ne peut pas retirer la pièce jointe", not fx["perms"]["delete"])
st, r = rd.req("DELETE", f"{base}/files/{f2['id']}")
check("lecteur : retrait refusé", st == 403, str(r))
st, link2 = lp.req("GET", f"{base}/files/{f2['id']}/link")
x = S()
st, data = x.req("GET", link2["path"] + "?inline=1")
check("aperçu PDF en ligne", st == 200 and "inline" in x.last_headers.get("content-disposition", ""), x.last_headers.get("content-disposition"))
st, r = lp.req("DELETE", f"{base}/files/{f2['id']}")
check("retrait par l'auteur", st == 200, str(r))
st, h = lp.req("GET", f"{base}/questions/{q2['id']}/history")
check("historique : ajout et retrait de pièce jointe", "attach" in [e["action"] for e in h] and "detach" in [e["action"] for e in h])

# ---------------------------------------------------------------------------
# Notifications
# ---------------------------------------------------------------------------
# par défaut, aucun e-mail : chacun l'active dans Mon compte
before = mark()
create(wf, f"Notification 0 {RUN}", ["Sécurité"], assignedParty="CLIENT")
time.sleep(1.5)
to = {m["to"] for m in outbox(before)}
check("par défaut : pas d'e-mail d'attribution", not (to & {LP, SEC, BOTH, RD}), str(to))
for s_ in (wf, lp, rd, sec, both):
    s_.req("POST", f"{base}/notify", {"notify": "IMMEDIATE"})
before = mark()
st, q8 = create(wf, f"Notification {RUN}", ["Sécurité"], assignedParty="CLIENT")
time.sleep(1.5)
new = outbox(before)
to = {m["to"] for m in new}
check("notification : éditeurs du client sur le stream", LP in to and SEC in to and BOTH in to, str(to))
check("notification : ni lecteur, ni auteur, ni Wifirst seul", RD not in to and WF not in to, str(to))
check("notification : objet et lien", all(f"n°{q8['ref']}" in m["subject"] and f"/c/la-poste?q={q8['ref']}" in m["text"] for m in new if m["to"] == LP), str([m["subject"] for m in new])[:300])
# préférences : aucun e-mail, récapitulatif
st, r = lp.req("POST", f"{base}/notify", {"notify": "DAILY"})
check("préférence enregistrée", st == 200, str(r))
sec.req("POST", f"{base}/notify", {"notify": "NONE"})
before = mark()
create(wf, f"Notification 2 {RUN}", ["Sécurité"], assignedParty="CLIENT")
time.sleep(1.5)
to = {m["to"] for m in outbox(before)}
check("préférence quotidienne ou aucune : pas d'e-mail immédiat", LP not in to and SEC not in to and BOTH in to, str(to))
before = mark()
st, r = S().req("POST", "/api/bridge/dev/digest", {})
time.sleep(0.5)
dig = [m for m in outbox(before) if m["to"] == LP]
check("récapitulatif quotidien", st == 200 and dig and "à traiter" in dig[0]["subject"], str(dig)[:300])
lp.req("POST", f"{base}/notify", {"notify": "IMMEDIATE"})
su.req("PATCH", f"/api/bridge/admin/clients/{CID}", {"settings": {"notifications": False}})
before = mark()
create(wf, f"Notification 3 {RUN}", ["Sécurité"], assignedParty="CLIENT")
time.sleep(1.5)
check("e-mails coupés pour le client", not [m for m in outbox(before) if m["to"] in (LP, BOTH)])
su.req("PATCH", f"/api/bridge/admin/clients/{CID}", {"settings": {"notifications": True}})
# réponse du client : Wifirst prévenu ; clôture : le créateur prévenu
before = mark()
lp.req("POST", f"{base}/questions/{q8['id']}/messages", {"body": "Réponse", "outcome": "PROVIDER"})
time.sleep(1.5)
to = {m["to"] for m in outbox(before)}
check("réponse attribuée à Wifirst : éditeurs Wifirst prévenus", WF in to and LP not in to, str(to))
before = mark()
wf.req("POST", f"{base}/questions/{q2['id']}/close", {})
time.sleep(1.5)
check("clôture : le créateur est prévenu", any(m["to"] == LP and "clôturée" in m["subject"] for m in outbox(before)))

# ---------------------------------------------------------------------------
# Recherche et export
# ---------------------------------------------------------------------------
st, r = lp.req("GET", f"{base}/search?q=" + urllib.parse.quote("CORSE"))
check("recherche dans les échanges, sans accents ni casse", st == 200 and q1["id"] in r["ids"], str(r))
st, r = sec.req("GET", f"{base}/search?q=" + urllib.parse.quote("corse"))
check("recherche : rien hors droits", st == 200 and q1["id"] not in r["ids"])
st, x = lp.req("GET", f"{base}/export.xlsx?status=all")
try:
    import openpyxl

    # fiche navette : un seul onglet, colonnes de réponse, identifiants cachés
    wb = openpyxl.load_workbook(io.BytesIO(x))
    check("fiche navette : un seul onglet", wb.sheetnames == ["Fiche navette"], str(wb.sheetnames))
    ws = wb["Fiche navette"]
    hdr = {c.value: c.column for c in ws[5] if c.value}
    check("fiche navette : colonnes", all(k in hdr for k in ["Réf.", "Sujet", "Question", "Échanges", "Votre réponse", "Nouvel attribué", "ID", "Version"]), str(hdr))
    check("fiche navette : identifiants cachés", ws.column_dimensions[openpyxl.utils.get_column_letter(hdr["ID"])].hidden)
    subjects = [r[hdr["Sujet"] - 1].value for r in ws.iter_rows(min_row=6)]
    check("fiche navette : questions visibles", any(s and s.startswith("Liste des ATM") for s in subjects), str(subjects)[:200])
    st, x2 = sec.req("GET", f"{base}/export.xlsx?status=all")
    w2 = openpyxl.load_workbook(io.BytesIO(x2))["Fiche navette"]
    s2 = [r[1].value for r in w2.iter_rows(min_row=6)]
    check("fiche navette : limitée aux droits", all("ATM" not in (s or "") for s in s2), str(s2)[:200])
    st, x3 = lp.req("GET", f"{base}/export.xlsx?status=open&assigned=CLIENT")
    w3 = openpyxl.load_workbook(io.BytesIO(x3))["Fiche navette"]
    a3 = {r[hdr["Attribuée à"] - 1].value for r in w3.iter_rows(min_row=6)}
    check("fiche navette : filtre d'attribution", a3 == {"La Poste"}, str(a3))

    def navette(session, query, fills):
        """Exporte la fiche navette, remplit les colonnes de réponse ({réf: (réponse, nouvel attribué)}) et renvoie le fichier."""
        st, data = session.req("GET", f"{base}/export.xlsx?{query}")
        book = openpyxl.load_workbook(io.BytesIO(data))
        sh = book["Fiche navette"]
        h = {c.value: c.column for c in sh[5] if c.value}
        for row in sh.iter_rows(min_row=6):
            ref = row[h["Réf."] - 1].value
            if ref in fills:
                sh.cell(row[0].row, h["Votre réponse"]).value = fills[ref][0]
                sh.cell(row[0].row, h["Nouvel attribué"]).value = fills[ref][1]
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()

    def preview(session, data):
        return session.req("POST", f"{base}/navette/preview", raw=data, headers={"Content-Type": "application/octet-stream"})

    def apply(session, items):
        return session.req("POST", f"{base}/navette/apply", {"items": [{k: i[k] for k in ("line", "questionId", "body", "target", "version")} for i in items if not i["error"]]})

    _, na = create(wf, f"Navette A {RUN}", ["Déploiement"], assignedParty="CLIENT")
    _, nb = create(wf, f"Navette B {RUN}", ["Déploiement"], assignedParty="CLIENT")
    _, nc = create(lp, f"Navette C {RUN}", ["Déploiement"])
    _, nd = create(wf, f"Navette D {RUN}", ["Déploiement"], assignedParty="CLIENT")
    _, ne = create(wf, f"Navette E {RUN}", ["Déploiement"], assignedParty="CLIENT")
    _, nf = create(wf, f"Navette F {RUN}", ["Déploiement"], assignedParty="CLIENT")
    fills = {
        na["ref"]: ("Voici la liste.\nAvec un retour à la ligne.", None),
        nb["ref"]: ("Premiers éléments, la suite vendredi.", "La Poste"),
        nc["ref"]: ("Je complète ma question.", None),
        nd["ref"]: (None, "Clôturer"),
        ne["ref"]: (None, "wifirst"),
        nf["ref"]: ("Réponse", "Peut-être"),
    }
    data = navette(lp, "status=open", fills)
    st, p = preview(lp, data)
    by = {i["ref"]: i for i in p["items"]} if st == 200 else {}
    check("navette, aperçu : lignes remplies seulement", st == 200 and set(by) == set(fills) and p["ready"] == 4, str(p)[:400])
    check("navette, aperçu : réponse de l'attribué, renvoyée à l'autre organisation", by.get(na["ref"], {}).get("action") == "answer" and by[na["ref"]]["outcome"] == "PROVIDER" and "Réponse de La Poste" in by[na["ref"]]["summary"], str(by.get(na["ref"]))[:300])
    check("navette, aperçu : attribution conservée", by.get(nb["ref"], {}).get("outcome") == "CLIENT", str(by.get(nb["ref"]))[:300])
    check("navette, aperçu : question attribuée à Wifirst refusée au client", "vous ne pouvez pas y répondre" in (by.get(nc["ref"], {}).get("error") or ""), str(by.get(nc["ref"]))[:300])
    check("navette, aperçu : clôture et changement d'attribution sans réponse", by.get(nd["ref"], {}).get("action") == "close" and by.get(ne["ref"], {}).get("action") == "assign", str([by.get(nd["ref"]), by.get(ne["ref"])])[:300])
    check("navette, aperçu : nouvel attribué non reconnu", "non reconnu" in (by.get(nf["ref"], {}).get("error") or ""), str(by.get(nf["ref"]))[:300])
    st, r = apply(lp, p["items"])
    check("navette : import appliqué", st == 200 and r["done"] == 4 and r["failed"] == 0, str(r)[:300])
    da = lp.req("GET", f"{base}/questions/{na['id']}")[1]
    check("navette : réponse enregistrée, marquée fiche navette", da["assignedParty"] == "PROVIDER" and da["status"] == "OPEN" and da["messages"][-1]["source"] == "navette" and da["messages"][-1]["party"] == "CLIENT" and da["messages"][-1]["body"] == "Voici la liste.\nAvec un retour à la ligne.", str(da)[:300])
    db_ = lp.req("GET", f"{base}/questions/{nb['id']}")[1]
    check("navette : réponse partielle, en cours", db_["status"] == "IN_PROGRESS" and db_["assignedParty"] == "CLIENT", str(db_)[:200])
    check("navette : clôture et attribution", lp.req("GET", f"{base}/questions/{nd['id']}")[1]["status"] == "CLOSED" and lp.req("GET", f"{base}/questions/{ne['id']}")[1]["assignedParty"] == "PROVIDER")
    st, h = lp.req("GET", f"{base}/questions/{na['id']}/history")
    check("navette : historique", any("fiche navette" in e["summary"] for e in h), str([e["summary"] for e in h])[:300])
    st, p2 = preview(lp, data)
    check("navette : second import du même fichier sans effet", st == 200 and p2["ready"] == 0 and "déjà dans les échanges" in (next(i for i in p2["items"] if i["ref"] == na["ref"])["error"] or ""), str(p2)[:400])
    # question clôturée : un nouvel attribué la rouvre avec la réponse
    data = navette(lp, "status=all", {nd["ref"]: ("Finalement, il manque un site.", "Wifirst")})
    st, p3 = preview(lp, data)
    i3 = p3["items"][0] if st == 200 and p3["items"] else {}
    check("navette : réouverture d'une question clôturée", i3.get("action") == "reopen", str(p3)[:300])
    apply(lp, p3["items"])
    dd = lp.req("GET", f"{base}/questions/{nd['id']}")[1]
    check("navette : rouverte, attribuée à Wifirst, avec la réponse", dd["status"] == "OPEN" and dd["assignedParty"] == "PROVIDER" and dd["messages"][-1]["outcome"] == "REOPEN", str(dd)[:300])
    # Wifirst complète à la place du client ; la question a changé depuis l'export : avertissement
    _, ng = create(wf, f"Navette G {RUN}", ["Déploiement"], assignedParty="CLIENT")
    data = navette(wf, "status=open", {ng["ref"]: ("Réponse transmise par La Poste en réunion.", None)})
    wf.req("PATCH", f"{base}/questions/{ng['id']}", {"dueDate": "2026-12-01"})
    st, p4 = preview(wf, data)
    i4 = p4["items"][0] if st == 200 and p4["items"] else {}
    check("navette, Wifirst : réponse de Wifirst à la place du client, attribution conservée", i4.get("party") == "PROVIDER" and i4.get("outcome") == "CLIENT", str(i4)[:300])
    check("navette : avertissement si la question a changé depuis l'export", "changé depuis l'export" in (i4.get("warning") or ""), str(i4)[:300])
    # fichiers refusés
    st, r = preview(lp, b"pas un fichier excel")
    check("navette : fichier illisible refusé", st == 400 and "illisible" in str(r), str(r)[:200])
    empty = io.BytesIO()
    openpyxl.Workbook().save(empty)
    st, r = preview(lp, empty.getvalue())
    check("navette : classeur sans colonnes de réponse refusé", st == 400 and "fiche navette" in str(r), str(r)[:200])
except ImportError:
    print("    (openpyxl absent : contrôle du fichier Excel ignoré)")

# ---------------------------------------------------------------------------
# Administration
# ---------------------------------------------------------------------------
st, r = lp.req("GET", "/api/bridge/admin/users")
check("administration réservée au super-admin", st == 403, str(r))
st, c = su.req("POST", "/api/bridge/admin/clients", {"name": f"Client test {RUN}", "clientName": "SNCF", "shortName": "SN", "streams": ["Déploiement", "Sécurité"]})
check("création d'un client avec ses streams", st == 200 and c["slug"].startswith("client-test"), str(c)[:200])
st, cl = su.req("GET", "/api/bridge/admin/clients")
cc = next(x for x in cl if x["id"] == c["id"])
check("client créé : streams et mode d'emploi", [s["name"] for s in cc["streams"]] == ["Déploiement", "Sécurité"] and "SNCF" in cc["description"], str(cc)[:300])
st, ns = su.req("POST", f"/api/bridge/admin/clients/{c['id']}/streams", {"name": "Gouvernance", "emoji": "🏛️"})
check("stream ajouté", st == 200 and ns["order"] == 3)
st, r = su.req("POST", f"/api/bridge/admin/clients/{c['id']}/streams/reorder", {"ids": [ns["id"], cc["streams"][0]["id"], cc["streams"][1]["id"]]})
check("streams réordonnés", st == 200)
st, r = su.req("PATCH", f"/api/bridge/admin/streams/{ns['id']}", {"name": "Gouvernance projet", "active": False})
check("stream renommé et désactivé", st == 200 and r["name"] == "Gouvernance projet" and r["active"] is False)
st, r = su.req("DELETE", f"/api/bridge/admin/streams/{ns['id']}")
check("stream inutilisé supprimé", st == 200)
st, r = su.req("DELETE", f"/api/bridge/admin/streams/{SID['Déploiement']}")
check("stream utilisé : suppression refusée", st == 400 and "désactivez" in r["error"], str(r))
st, r = lp.req("GET", f"/api/bridge/c/{c['slug']}")
check("client non ouvert : accès refusé", st == 403, str(r))
# droits : accès à plusieurs clients
lp_user = next(u for u in su.req("GET", "/api/bridge/admin/users")[1] if u["email"] == LP)
st, r = su.req("PUT", f"/api/bridge/admin/users/{lp_user['id']}/memberships", {"memberships": [{"clientId": CID, "side": "CLIENT", "defaultAccess": "CLIENT", "streamAccess": {SID["Gouvernance"]: "READ", SID["Sécurité"]: "CLIENT"}}, {"clientId": c["id"], "side": "CLIENT", "defaultAccess": "READ", "streamAccess": {}}]})
check("droits sur deux clients", st == 200, str(r))
st, me = lp.req("GET", "/api/bridge/auth/me")
check("deux clients ouverts", {x["slug"] for x in me["clients"]} == {"la-poste", c["slug"]}, str(me)[:300])
st, bb = lp.req("GET", base)
check("droit par stream : lecture sur Gouvernance, défaut ailleurs (identique au défaut non conservé)", bb["me"]["access"][SID["Gouvernance"]] == "READ" and bb["me"]["access"][SID["Déploiement"]] == "CLIENT")
lpm = next(u for u in su.req("GET", "/api/bridge/admin/users")[1] if u["email"] == LP)["memberships"]
check("droit identique au défaut non conservé", all(SID["Sécurité"] not in m["streamAccess"] for m in lpm), str(lpm)[:300])
su.req("PATCH", f"/api/bridge/admin/clients/{c['id']}", {"archived": True})
st, r = lp.req("GET", f"/api/bridge/c/{c['slug']}")
check("client archivé : plus ouvert aux membres", st == 403, str(r))

# retrait de l'accès WiBridge : session refusée dès la requête suivante
rd_user = next(u for u in su.req("GET", "/api/bridge/admin/users")[1] if u["email"] == RD)
su.req("PATCH", f"/api/bridge/admin/users/{rd_user['id']}", {"bridgeAccess": False})
check("accès WiBridge retiré : session refusée", rd.req("GET", "/api/bridge/auth/me")[0] == 401)
check("accès WiBridge retiré : connexion refusée", S().req("POST", "/api/bridge/auth/login", {"email": RD, "password": UPWD})[0] == 401)
st, r = su.req("PATCH", f"/api/bridge/admin/users/{next(u for u in su.req('GET', '/api/bridge/admin/users')[1] if u['isSuperAdmin'])['id']}", {"bridgeAccess": False})
check("super-admin : accès conservé", st == 400, str(r))

# mot de passe oublié : nouveau mot de passe, appareils oubliés
x = S()
before = mark()
st, r = x.req("POST", "/api/bridge/auth/forgot", {"email": BOTH})
check("mot de passe oublié : code", st == 200 and r.get("devCode"), str(r))
check("mot de passe oublié : e-mail du code envoyé", any(m["to"] == BOTH and r["devCode"] in m["subject"] for m in outbox(before)))
st, r2 = x.req("POST", "/api/bridge/auth/reset", {"challengeId": r["challengeId"], "code": r["devCode"], "password": "Bridge-Nouveau-2026"})
check("nouveau mot de passe", st == 200, str(r2))
check("ancienne session fermée", both.req("GET", "/api/bridge/auth/me")[0] == 401)
check("appareils oubliés : code redemandé", blogin(both, BOTH, "Bridge-Nouveau-2026") == "code")
# changement de mot de passe : la session en cours reste ouverte
st, r = wf.req("POST", "/api/bridge/auth/password", {"current": UPWD, "next": "Bridge-Change-2026"})
check("changement de mot de passe", st == 200, str(r))
check("changement : session en cours conservée", wf.req("GET", "/api/bridge/auth/me")[0] == 200)
st, devs = wf.req("GET", "/api/bridge/auth/devices")
check("appareils : celui-ci reste de confiance", st == 200 and any(d["current"] for d in devs), str(devs))
check("changement : connexion sans code sur cet appareil", blogin(wf, WF, "Bridge-Change-2026") == "trusted")
st, r = wf.req("POST", "/api/bridge/auth/devices/revoke-all", {})
check("tous les appareils retirés", st == 200 and blogin(wf, WF, "Bridge-Change-2026") == "code")

# journal
st, j = su.req("GET", f"/api/bridge/admin/journal?clientId={CID}&limit=50")
check("journal du client", st == 200 and len(j) > 10 and j[0]["clientName"] == "La Poste", str(j[:1])[:200])

print()
print("RESULT:", "ALL OK" if not fails else f"{len(fails)} FAIL : {fails}")
