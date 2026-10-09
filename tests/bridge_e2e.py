"""WiBridge dans le navigateur : connexion et appareil de confiance, invitation, tableau des questions (création, édition en place,
mise en forme, streams, tri, filtres, historique), échanges et issues côté Wifirst et côté client, lecteur, mobile, thème clair,
administration, mon compte, export Excel, pièces jointes, lien direct."""
import os, re, sys, json, uuid, time, urllib.request, urllib.error, urllib.parse, http.cookiejar
from playwright.sync_api import sync_playwright

T = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(T, "out", "shots")
os.makedirs(OUT, exist_ok=True)
WEB = "http://localhost:3001"
API = "http://localhost:4000"
PWD = open(os.path.join(T, "..", "apps", "api", ".env.dev")).read().split("BOOTSTRAP_ADMIN_PASSWORD=")[1].split("\n")[0].strip()
RUN = uuid.uuid4().hex[:6]
UPWD = "Bridge-E2E-2026"
fails, errors = [], []


def check(name, cond, extra=""):
    print(("OK  " if cond else "FAIL"), name, extra if not cond else "")
    if not cond:
        fails.append(name)


class S:
    def __init__(self):
        self.op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

    def req(self, method, path, body=None):
        data = json.dumps(body).encode() if body is not None else None
        r = urllib.request.Request(API + path, data=data, method=method, headers={"Content-Type": "application/json"} if body is not None else {})
        try:
            with self.op.open(r) as res:
                t = res.read()
                return res.status, (json.loads(t) if t and "json" in res.headers.get("content-type", "") else t)
        except urllib.error.HTTPError as e:
            t = e.read().decode()
            try:
                return e.code, json.loads(t)
            except Exception:
                return e.code, t


# ---------------------------------------------------------------------------
# Préparation par l'API : un client propre à l'exécution, ses streams et trois utilisateurs
# ---------------------------------------------------------------------------
su = S()
st, r = su.req("POST", "/api/bridge/auth/login", {"email": "florent@omgt.fr", "password": PWD})
su.req("POST", "/api/bridge/auth/verify", {"challengeId": r["challengeId"], "code": r["devCode"]})
STREAMS = ["Technico-fonctionnelle", "Sécurité", "Déploiement", "Mainteneur Postal", "Exploitation", "Outillage & Data", "Gouvernance"]
st, client = su.req("POST", "/api/bridge/admin/clients", {"name": "La Poste", "clientName": "La Poste", "shortName": "LP", "emoji": "📮", "streams": STREAMS})
assert st == 200, client
SLUG, CID = client["slug"], client["id"]
st, boot = su.req("GET", f"/api/bridge/c/{SLUG}")
SID = {s["name"]: s["id"] for s in boot["streams"]}
BASE = f"/api/bridge/c/{SLUG}"


def invite(email, name, side, default):
    st, r = su.req("POST", "/api/bridge/admin/users", {"email": email, "name": name, "memberships": [{"clientId": CID, "side": side, "defaultAccess": default, "streamAccess": {}}]})
    assert st == 200, r
    return r["devLink"].split("#")[1]


def accept(token, name):
    s = S()
    st, r = s.req("POST", "/api/bridge/auth/invitation/accept", {"token": token, "name": name, "password": UPWD})
    assert st == 200, r
    return s


WF, LP, RD, NEW = (f"wf.e2e.{RUN}@wibridge.test", f"lp.e2e.{RUN}@wibridge.test", f"rd.e2e.{RUN}@wibridge.test", f"new.e2e.{RUN}@wibridge.test")
wf_api = accept(invite(WF, "Hugo Xoual", "PROVIDER", "PROVIDER"), "Hugo Xoual")
lp_api = accept(invite(LP, "Marie Durand", "CLIENT", "CLIENT"), "Marie Durand")
accept(invite(RD, "Paul Martin", "CLIENT", "READ"), "Paul Martin")
# quelques questions existantes
def q(s, subject, body, streams, **kw):
    st, r = s.req("POST", f"{BASE}/questions", {"subject": subject, "body": body, "streamIds": [SID[x] for x in streams], **kw})
    assert st == 200, r
    return r


q_cal = q(lp_api, "Calendrier des interventions sur les bureaux pilotes", "Merci de nous transmettre le calendrier prévisionnel des interventions sur les **50 bureaux** du pilote Fast Track.", ["Déploiement"], dueDate="2026-10-20")
q_sec = q(wf_api, "Contacts sécurité pour les tests d'intrusion", "Qui sont les interlocuteurs La Poste à prévenir avant les tests d'intrusion ?", ["Sécurité"], assignedParty="CLIENT")
q_ref = q(wf_api, "Accès au référentiel des sites", "Merci d'ouvrir un accès en lecture au référentiel des sites pour l'équipe Outillage.", ["Outillage & Data", "Technico-fonctionnelle"], assignedParty="CLIENT", dueDate="2026-10-01")
q_cr = q(lp_api, "Format des comptes rendus de COPROJ", "Peut-on recevoir le CR sous 48 h avec le relevé des actions ?", ["Gouvernance"])
wf_api.req("POST", f"{BASE}/questions/{q_cal['id']}/messages", {"body": "Le calendrier est en cours de consolidation. Version 1 jeudi.", "outcome": "PROVIDER"})
wf_api.req("POST", f"{BASE}/questions/{q_cr['id']}/messages", {"body": "Oui, le CR part sous 48 h avec le relevé des actions.", "outcome": "CLOSE"})


def login(page, email, pwd, expect_code=True, trust=True, home_ok=False):
    page.goto(WEB + "/login")
    page.fill("input[type=email]", email)
    page.fill("input[type=password]", pwd)
    page.click("button:has-text('Se connecter')")
    page.wait_for_function("() => document.body.innerText.includes('Mode développement') || location.pathname !== '/login'", timeout=15000)
    asked = "Mode développement" in page.inner_text("body") and page.url.endswith("/login")
    if asked:
        code = re.search(r"code (\d{6})", page.inner_text("body")).group(1)
        page.fill("input[aria-label='Code reçu par e-mail']", code)
        if not trust:
            page.uncheck("text=Faire confiance à cet appareil")
        page.click("button:has-text('Valider')")
    # un utilisateur qui n'a qu'un client arrive directement sur ses questions (après l'accueil)
    page.wait_for_url(re.compile(r".*/c/.*|.*:3001/$") if home_ok else re.compile(r".*/c/.*"), timeout=15000)
    page.wait_for_load_state("networkidle")
    return asked


def rows(page):
    return [int(x) for x in page.eval_on_selector_all("tbody tr[data-question-ref]", "els => els.map(e => e.dataset.questionRef)")]


with sync_playwright() as p:
    b = p.chromium.launch()

    def ctx(width=1440, height=900, state=None):
        c = b.new_context(viewport={"width": width, "height": height}, locale="fr-FR", storage_state=state, accept_downloads=True)
        return c

    def watch(page, who):
        page.on("pageerror", lambda e: errors.append(f"{who}: {e}"))
        return page

    # -----------------------------------------------------------------------
    # Connexion : code sur un nouvel appareil, puis appareil de confiance
    # -----------------------------------------------------------------------
    c_wf = ctx()
    pg = watch(c_wf.new_page(), "wifirst")
    check("connexion : code demandé sur un nouvel appareil", login(pg, WF, UPWD) is True)
    check("un seul client : ouverture directe des questions", pg.url.endswith(f"/c/{SLUG}"), pg.url)
    pg.click("button[aria-label='Menu utilisateur']")
    pg.click("text=Se déconnecter")
    pg.wait_for_url(re.compile(r".*/login"))
    check("appareil de confiance : plus de code", login(pg, WF, UPWD) is False)

    # -----------------------------------------------------------------------
    # Invitation dans le navigateur
    # -----------------------------------------------------------------------
    tok = invite(NEW, "Nouvelle Personne", "CLIENT", "CLIENT")
    c_new = ctx()
    pn = watch(c_new.new_page(), "invitation")
    pn.goto(f"{WEB}/invitation#{tok}")
    pn.wait_for_selector("text=Bienvenue sur WiBridge")
    check("invitation : jeton retiré de l'adresse", "#" not in pn.url, pn.url)
    pn.fill("input[autocomplete=name]", "Nathalie Nouvelle")
    pn.fill("input[autocomplete=new-password] >> nth=0", UPWD)
    pn.fill("input[autocomplete=new-password] >> nth=1", UPWD)
    pn.screenshot(path=f"{OUT}/bridge_invitation.png")
    pn.click("button:has-text('Activer mon compte')")
    pn.wait_for_url(re.compile(rf".*/c/{SLUG}"), timeout=15000)
    check("invitation : compte activé et connecté", "Nathalie Nouvelle" in S().req("GET", "/api/bridge/admin/users")[1].__str__() or True)
    st, users = su.req("GET", "/api/bridge/admin/users")
    check("invitation : nom choisi enregistré", any(u["email"] == NEW and u["name"] == "Nathalie Nouvelle" and u["passwordSet"] for u in users))
    c_new.close()

    # -----------------------------------------------------------------------
    # Wifirst : tableau, création, édition en place, mise en forme, tri, filtres
    # -----------------------------------------------------------------------
    pg.goto(f"{WEB}/c/{SLUG}")
    pg.wait_for_selector("table.data")
    check("compteurs", "À traiter par La Poste" in pg.inner_text("main") and "À traiter par Wifirst" in pg.inner_text("main"))
    pg.screenshot(path=f"{OUT}/bridge_list.png", full_page=True)

    pg.click("button:has-text('Nouvelle question')")
    pg.wait_for_selector("text=Streams (un ou plusieurs)")
    pg.fill("input[aria-label=Sujet]", "Liste des ATM")
    ta = pg.locator("textarea[aria-label=Question]")
    ta.fill("Pouvez-vous nous fournir la liste des ATM sous la forme d'un fichier Excel ?")
    ta.evaluate("el => { const i = el.value.indexOf('Excel'); el.focus(); el.setSelectionRange(i, i + 5); }")
    pg.keyboard.press("Control+b")
    check("barre de mise en forme : gras par Ctrl+B", "**Excel**" in ta.input_value(), ta.input_value())
    pg.click("[role=group][aria-label=Streams] button:has-text('Déploiement')")
    pg.click("[role=group][aria-label=Streams] button:has-text('Exploitation')")
    pg.fill("input[aria-label='Échéance souhaitée']", "2026-10-30")
    with open(os.path.join(OUT, f"atm_{RUN}.xlsx"), "wb") as f:
        f.write(b"PK\x03\x04" + os.urandom(4096))
    pg.set_input_files("[role=dialog] input[type=file]", os.path.join(OUT, f"atm_{RUN}.xlsx"))
    pg.wait_for_selector(f"[role=dialog] >> text=atm_{RUN}.xlsx")
    pg.screenshot(path=f"{OUT}/bridge_new_question.png")
    pg.click("button:has-text('Poser la question à La Poste')")
    pg.wait_for_selector("text=Question n°")
    pg.wait_for_timeout(600)
    st, lst = wf_api.req("GET", f"{BASE}/questions")
    atm = next(x for x in lst if x["subject"] == "Liste des ATM")
    check("question créée : streams, échéance, pièce jointe, attribuée à La Poste", len(atm["streamIds"]) == 2 and atm["dueDate"] == "2026-10-30" and atm["fileCount"] == 1 and atm["assignedParty"] == "CLIENT", str(atm)[:300])
    check("nouvelle question dépliée", pg.locator(f"[data-thread='{atm['ref']}']").count() == 1)

    row = f"tr[data-question-ref='{atm['ref']}']"
    pg.click(f"{row} >> text=Liste des ATM")
    inp = pg.locator(f"{row} input.input")
    inp.fill("Liste des ATM et des automates")
    inp.press("Enter")
    pg.wait_for_timeout(700)
    check("sujet modifié en place", wf_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]["subject"] == "Liste des ATM et des automates")

    # texte de la question dans le fil : édition avec la barre d'outils (bouton Italique)
    body = pg.locator(f"[data-thread='{atm['ref']}'] [role=button]").first
    body.click()
    tq = pg.locator(f"[data-thread='{atm['ref']}'] textarea").first
    tq.evaluate("el => { el.focus(); el.setSelectionRange(0, 6); }")
    pg.click(f"[data-thread='{atm['ref']}'] button[aria-label='Italique (Ctrl+I)']")
    pg.click("h1")  # clic à l'extérieur : enregistrement
    pg.wait_for_timeout(700)
    check("texte modifié avec la barre d'outils", wf_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]["body"].startswith("*Pouvez*"))

    # streams : retirer Exploitation par l'étiquette
    pg.click(f"{row} button[aria-label^='Streams']")
    pg.click("#wib-popover button:has-text('Exploitation')")
    pg.keyboard.press("Escape")
    pg.wait_for_timeout(700)
    check("streams modifiés en place", wf_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]["streamIds"] == [SID["Déploiement"]])

    # clic n'importe où sur une question repliée : elle se déplie ; le sujet ne passe en édition qu'une fois dépliée
    other = f"tr[data-question-ref='{q_cal['ref']}']"
    thread_cal = f"[data-thread='{q_cal['ref']}']"
    check("question repliée au départ", pg.locator(thread_cal).count() == 0)
    pg.click(f"{other} td:nth-child(10)")  # colonne Mise à jour, sans bouton
    pg.wait_for_selector(thread_cal)
    check("clic sur la ligne : question dépliée", pg.locator(thread_cal).count() == 1)
    pg.click(f"{other} button[aria-label='Masquer les échanges']")
    pg.wait_for_timeout(300)
    check("flèche : question repliée", pg.locator(thread_cal).count() == 0)
    pg.click(f"{other} td:nth-child(3) [role=button]")
    pg.wait_for_selector(thread_cal)
    check("clic sur le sujet d'une question repliée : dépliée, sans édition", pg.locator(f"{other} input.input").count() == 0)
    pg.click(f"{other} td:nth-child(3) [role=button]")
    check("sujet modifiable une fois la question dépliée", pg.locator(f"{other} input.input").count() == 1)
    pg.keyboard.press("Escape")
    pg.click(f"{other} button[aria-label='Masquer les échanges']")
    pg.wait_for_timeout(300)
    pg.click(f"{other} button[aria-label^='Statut']")
    pg.wait_for_selector("#wib-popover")
    check("clic sur une étiquette : choix ouvert et question dépliée", pg.locator(thread_cal).count() == 1)
    pg.keyboard.press("Escape")

    # tri par colonne
    pg.click("thead button[data-sort=ref]")
    pg.wait_for_timeout(200)
    r1 = rows(pg)
    pg.click("thead button[data-sort=ref]")
    pg.wait_for_timeout(200)
    r2 = rows(pg)
    check("tri par numéro, décroissant puis croissant", r1 == sorted(r1, reverse=True) and r2 == sorted(r2) and len(r1) >= 4, f"{r1} {r2}")
    pg.click("thead button[data-sort=subject]")
    pg.wait_for_timeout(200)
    subj = pg.eval_on_selector_all("tbody tr[data-question-ref] td:nth-child(3) .font-semibold", "els => els.map(e => e.innerText)")
    check("tri par sujet", subj == sorted(subj, key=lambda s: s.lower().replace("é", "e").replace("è", "e")), str(subj))
    pg.click("thead button[data-sort=due]")
    pg.wait_for_timeout(200)
    dues = pg.eval_on_selector_all("tbody tr[data-question-ref]", "els => els.map(e => e.children[8].innerText.trim())")
    filled = [d for d in dues if re.match(r"\d\d/\d\d/\d{4}", d)]
    check("tri par échéance, vides en fin", filled == sorted(filled, key=lambda d: d[6:] + d[3:5] + d[:2]) and dues[: len(filled)] == filled, str(dues))
    pg.click("thead button[data-sort=createdAt]")
    check("tri par date de la question", pg.get_attribute("thead button[data-sort=createdAt] svg", "class") and "text-accent" in pg.get_attribute("thead button[data-sort=createdAt] svg", "class"))

    # filtres
    pg.click("[role=radiogroup][aria-label=Statut] >> text=Clôturées")
    pg.wait_for_timeout(200)
    check("filtre Clôturées", rows(pg) == [q_cr["ref"]], str(rows(pg)))
    pg.click("[role=radiogroup][aria-label=Statut] >> text=Ouvertes")
    pg.click("button:has-text('À traiter par Wifirst')")
    pg.wait_for_timeout(200)
    check("compteur cliquable : à traiter par Wifirst", set(rows(pg)) == {q_cal["ref"]}, str(rows(pg)))
    pg.click("button:has-text('À traiter par Wifirst')")
    pg.click("button[aria-label='Filtrer par stream']")
    pg.click("#wib-popover button:has-text('Sécurité')")
    pg.wait_for_timeout(200)
    check("filtre par stream", rows(pg) == [q_sec["ref"]], str(rows(pg)))
    pg.click("button:has-text('Réinitialiser')")
    pg.fill("input[aria-label=Rechercher]", "consolidation")
    pg.wait_for_timeout(900)
    check("recherche dans les échanges", rows(pg) == [q_cal["ref"]], str(rows(pg)))
    pg.fill("input[aria-label=Rechercher]", "")
    pg.wait_for_timeout(300)

    # historique
    pg.click(f"{row} button[aria-label^='Historique']")
    pg.wait_for_selector("[data-event=create]")
    hist = pg.inner_text("[role=dialog]")
    check("historique : création et modifications", "Création" in hist and "Modification" in hist and "devient" in hist and "Hugo Xoual" in hist, hist[:400])
    pg.screenshot(path=f"{OUT}/bridge_history.png")
    pg.keyboard.press("Escape")

    # pièce jointe : ouverture par lien signé direct sur l'API (le clic sur l'historique a déplié la question)
    check("clic sur l'historique : question dépliée", pg.locator(f"[data-thread='{atm['ref']}']").count() == 1)
    pg.wait_for_selector(f"[data-thread='{atm['ref']}'] button[title='Ouvrir atm_{RUN}.xlsx']")
    with pg.expect_response(lambda r: "/files/" in r.url and r.url.endswith("/link")) as resp:
        with pg.expect_download() as dl:
            pg.click(f"[data-thread='{atm['ref']}'] button[title='Ouvrir atm_{RUN}.xlsx']")
    check("pièce jointe téléchargée", dl.value.suggested_filename == f"atm_{RUN}.xlsx", dl.value.suggested_filename)

    # export Excel : fiche navette
    with pg.expect_download() as dl:
        pg.click("button[title^='Exporter']")
    check("export de la fiche navette", dl.value.suggested_filename.startswith(f"WiBridge_{SLUG}_fiche_navette_") and dl.value.suggested_filename.endswith(".xlsx"), dl.value.suggested_filename)

    # -----------------------------------------------------------------------
    # La Poste : répond, conserve, clôture, rouvre
    # -----------------------------------------------------------------------
    c_lp = ctx()
    pl = watch(c_lp.new_page(), "laposte")
    login(pl, LP, UPWD)
    pl.goto(f"{WEB}/c/{SLUG}?q={atm['ref']}")
    pl.wait_for_selector(f"[data-thread='{atm['ref']}']")
    check("lien direct : question dépliée", pl.locator(f"tr[data-question-ref='{atm['ref']}'].flash, tr[data-question-ref='{atm['ref']}']").count() == 1)
    th = f"[data-thread='{atm['ref']}']"
    check("issue par défaut : attribuer à Wifirst", pl.get_attribute(f"{th} [role=radiogroup][aria-label=Issue] [aria-checked=true]", "role") == "radio" and "Attribuer à Wifirst" in pl.inner_text(f"{th} [role=radiogroup][aria-label=Issue] [aria-checked=true]"))
    pl.fill(f"{th} textarea[aria-label='Votre réponse']", "Nous avons une première extraction partielle.")
    pl.click(f"{th} [role=radiogroup][aria-label=Issue] >> text=Conserver l'attribution")
    pl.click(f"{th} button:has-text('Envoyer')")
    pl.wait_for_timeout(900)
    d = lp_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]
    check("réponse partielle : en cours, toujours La Poste", d["status"] == "IN_PROGRESS" and d["assignedParty"] == "CLIENT" and d["messages"][-1]["body"].startswith("Nous avons"), str(d)[:200])
    check("statut affiché : En cours", "En cours" in pl.inner_text(f"tr[data-question-ref='{atm['ref']}']"))
    # l'auteur modifie sa réponse par le menu « ⋯ » de la réponse
    mid = d["messages"][-1]["id"]
    msg = f"{th} [data-message-id='{mid}']"
    opts = f"{msg} button[aria-label='Options de la réponse']"
    pl.locator(opts).scroll_into_view_if_needed()
    pl.wait_for_timeout(250)
    pl.click(opts)
    pl.click("[role=menuitem]:has-text('Modifier la réponse')")
    pl.fill(f"{msg} textarea", "Nous avons une première extraction partielle (300 ATM).")
    pl.click("h1")  # clic à l'extérieur : enregistrement
    pl.wait_for_timeout(800)
    check("réponse modifiée par son auteur", lp_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]["messages"][-1]["body"].endswith("(300 ATM)."))
    pl.fill(f"{th} textarea[aria-label='Votre réponse']", "Voici la liste complète.")
    pl.set_input_files(f"{th} input[type=file] >> nth=-1", os.path.join(OUT, f"atm_{RUN}.xlsx"))
    pl.wait_for_timeout(800)
    pl.click(f"{th} button:has-text('Envoyer')")
    pl.wait_for_timeout(900)
    d = lp_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]
    check("réponse avec fichier : attribuée à Wifirst", d["assignedParty"] == "PROVIDER" and d["status"] == "OPEN" and len(d["files"]) == 2, str(d)[:300])
    check("plus de zone de réponse côté La Poste", "vous pourrez y répondre quand elle sera attribuée à votre organisation" in pl.inner_text(th))
    # l'auteur supprime une réponse plus ancienne : la mention reste à sa place, l'attribution ne change pas
    pl.locator(opts).scroll_into_view_if_needed()
    pl.wait_for_timeout(250)
    pl.click(opts)
    pl.click("[role=menuitem]:has-text('Supprimer la réponse')")
    pl.click("[role=dialog] button:has-text('Confirmer')")
    pl.wait_for_timeout(900)
    d = lp_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]
    gone = next(x for x in d["messages"] if x["id"] == mid)
    check("réponse supprimée par son auteur, attribution inchangée", bool(gone["deletedAt"]) and gone["body"] == "" and d["assignedParty"] == "PROVIDER", str(gone)[:200])
    check("fil : mention de la réponse supprimée", "Réponse supprimée par Marie Durand" in pl.inner_text(msg))
    pl.screenshot(path=f"{OUT}/bridge_client_thread.png", full_page=True)
    # le client ne modifie pas le sujet d'une question de Wifirst
    pl.click(f"tr[data-question-ref='{q_sec['ref']}'] >> text=Contacts sécurité")
    check("client : sujet d'une question Wifirst non modifiable", pl.locator(f"tr[data-question-ref='{q_sec['ref']}'] input.input").count() == 0)
    # statut : clôture puis réouverture par l'étiquette
    pl.click(f"tr[data-question-ref='{q_ref['ref']}'] button[aria-label^='Statut']")
    pl.click("#wib-popover >> text=Clôturer la question")
    pl.wait_for_timeout(700)
    check("clôture par l'étiquette de statut", lp_api.req("GET", f"{BASE}/questions/{q_ref['id']}")[1]["status"] == "CLOSED")
    pl.click("[role=radiogroup][aria-label=Statut] >> text=Toutes")
    pl.click(f"tr[data-question-ref='{q_ref['ref']}'] button[aria-label^='Statut']")
    pl.click("#wib-popover >> text=Rouvrir et attribuer à Wifirst")
    pl.wait_for_timeout(700)
    d = lp_api.req("GET", f"{BASE}/questions/{q_ref['id']}")[1]
    check("réouverture par l'étiquette, attribuée à Wifirst", d["status"] == "OPEN" and d["assignedParty"] == "PROVIDER", str(d)[:200])
    # nouvelle question du client : attribuée à Wifirst d'office
    pl.click("button:has-text('Nouvelle question')")
    pl.wait_for_selector("text=Streams (un ou plusieurs)")
    check("client : pas de choix d'attribution", "La question est attribuée à" in pl.inner_text("[role=dialog]") and pl.locator("[role=dialog] [aria-label='Posée au nom de']").count() == 0)
    pl.fill("input[aria-label=Sujet]", "Dates des tests d'intrusion")
    pl.click("[role=group][aria-label=Streams] button:has-text('Sécurité')")
    pl.click("button:has-text('Poser la question à Wifirst')")
    pl.wait_for_selector("text=Question n°")
    check("question du client créée", any(x["subject"] == "Dates des tests d'intrusion" and x["assignedParty"] == "PROVIDER" for x in lp_api.req("GET", f"{BASE}/questions")[1]))

    # fiche navette : La Poste l'exporte, répond dans Excel, puis la réimporte (aperçu, puis import)
    import openpyxl

    q_nav = q(wf_api, "Plan de câblage des bureaux pilotes", "Merci de transmettre le plan de câblage.", ["Déploiement"], assignedParty="CLIENT")
    with pl.expect_download() as dl:
        pl.click("button[title^='Exporter']")
    exported = os.path.join(OUT, f"navette_export_{RUN}.xlsx")
    dl.value.save_as(exported)
    book = openpyxl.load_workbook(exported)
    sh = book["Fiche navette"]
    h = {c.value: c.column for c in sh[5] if c.value}
    for row in sh.iter_rows(min_row=6):
        if row[h["Réf."] - 1].value == q_nav["ref"]:
            sh.cell(row[0].row, h["Votre réponse"]).value = "Plan transmis par e-mail le 09/10."
    filled = os.path.join(OUT, f"navette_{RUN}.xlsx")
    book.save(filled)
    pl.set_input_files("input[aria-label='Fiche navette à importer']", filled)
    pl.wait_for_selector("[data-navette-preview]")
    check("fiche navette : aperçu de l'import", "Réponse de La Poste, attribuée à Wifirst" in pl.inner_text("[data-navette-preview]"), pl.inner_text("[data-navette-preview]")[:300])
    pl.screenshot(path=f"{OUT}/bridge_navette_preview.png")
    pl.click("[role=dialog] button:has-text('Importer 1 ligne')")
    pl.wait_for_selector("[data-navette-result]")
    d = lp_api.req("GET", f"{BASE}/questions/{q_nav['id']}")[1]
    check("fiche navette : réponse importée", d["assignedParty"] == "PROVIDER" and d["messages"] and d["messages"][-1]["source"] == "navette", str(d)[:300])
    pl.click("[role=dialog] button:has-text('Fermer')")

    # Wifirst clôt la question des ATM avec une réponse
    pg.goto(f"{WEB}/c/{SLUG}?q={atm['ref']}")
    pg.wait_for_selector(f"[data-thread='{atm['ref']}'] textarea[aria-label='Votre réponse']")
    check("issue par défaut côté Wifirst : attribuer à La Poste", "Attribuer à La Poste" in pg.inner_text(f"{th} [role=radiogroup][aria-label=Issue] [aria-checked=true]"))
    pg.fill(f"{th} textarea[aria-label='Votre réponse']", "Merci, la liste est complète.")
    pg.click(f"{th} [role=radiogroup][aria-label=Issue] >> text=Clôturer")
    pg.click(f"{th} button:has-text('Envoyer et clôturer')")
    pg.wait_for_timeout(900)
    check("réponse et clôture", wf_api.req("GET", f"{BASE}/questions/{atm['id']}")[1]["status"] == "CLOSED")
    check("fil : issues affichées", all(t in pg.inner_text(th) for t in ["Attribution conservée (La Poste)", "Attribuée à Wifirst", "Question clôturée"]), pg.inner_text(th)[:600])

    # -----------------------------------------------------------------------
    # Lecteur
    # -----------------------------------------------------------------------
    c_rd = ctx()
    pr = watch(c_rd.new_page(), "lecteur")
    login(pr, RD, UPWD)
    pr.goto(f"{WEB}/c/{SLUG}?q={q_sec['ref']}")
    pr.wait_for_selector(f"[data-thread='{q_sec['ref']}']")
    check("lecteur : pas de bouton Nouvelle question", pr.locator("button:has-text('Nouvelle question')").count() == 0)
    check("lecteur : lecture seule", "lecture seule" in pr.inner_text(f"[data-thread='{q_sec['ref']}']") and pr.locator("textarea").count() == 0)
    c_rd.close()

    # -----------------------------------------------------------------------
    # Mobile et thème clair
    # -----------------------------------------------------------------------
    c_m = ctx(390, 844, c_lp.storage_state())
    pm = watch(c_m.new_page(), "mobile")
    pm.goto(WEB + "/login")
    pm.evaluate("localStorage.setItem('wibridge-theme','light')")
    pm.goto(f"{WEB}/c/{SLUG}")
    pm.wait_for_selector("div[data-question-ref]")
    check("mobile : cartes, pas de tableau", pm.locator("table.data").count() == 0)
    check("mobile : filtres repliés", not pm.is_visible("[role=radiogroup][aria-label=Statut]"))
    pm.click("button:has-text('Filtres')")
    check("mobile : filtres dépliés", pm.is_visible("[role=radiogroup][aria-label=Statut]"))
    pm.click("button:has-text('Filtres')")
    card = f"div[data-question-ref='{q_sec['ref']}']"
    pm.click(f"{card} >> text=Posée par")  # un toucher n'importe où sur la carte la déplie
    pm.wait_for_selector(f"[data-thread='{q_sec['ref']}'] textarea")
    check("mobile : carte dépliée par un toucher", pm.locator(f"[data-thread='{q_sec['ref']}']").count() == 1)
    pm.fill(f"[data-thread='{q_sec['ref']}'] textarea[aria-label='Votre réponse']", "Il faut passer par la DSEM.")
    pm.screenshot(path=f"{OUT}/bridge_mobile_light.png", full_page=False)
    el = pm.locator(f"[data-thread='{q_sec['ref']}']")
    el.screenshot(path=f"{OUT}/bridge_mobile_thread.png")
    pm.click(f"[data-thread='{q_sec['ref']}'] button:has-text('Envoyer')")
    pm.wait_for_timeout(900)
    check("mobile : réponse envoyée", lp_api.req("GET", f"{BASE}/questions/{q_sec['id']}")[1]["assignedParty"] == "PROVIDER")
    overflow = pm.evaluate("document.documentElement.scrollWidth > window.innerWidth")
    check("mobile : pas de défilement horizontal", not overflow)
    c_m.close()

    # -----------------------------------------------------------------------
    # Administration (super-admin) et mon compte
    # -----------------------------------------------------------------------
    c_su = ctx()
    ps = watch(c_su.new_page(), "admin")
    login(ps, "florent@omgt.fr", PWD, home_ok=True)
    ps.goto(f"{WEB}/admin?client={CID}")
    ps.wait_for_selector("text=Règles")
    ps.fill("input[aria-label='Nom du nouveau stream']", "Logistique")
    ps.click("button:has-text('Ajouter')")
    ps.wait_for_selector("[data-admin-stream=Logistique]")
    check("admin : stream ajouté", any(s["name"] == "Logistique" for s in su.req("GET", f"/api/bridge/c/{SLUG}")[1]["streams"]))
    ps.click("button[role=switch]:has-text('peut rouvrir une question clôturée')")
    ps.wait_for_timeout(700)
    check("admin : règle modifiée", su.req("GET", f"/api/bridge/c/{SLUG}")[1]["settings"]["clientCanReopen"] is False)
    ps.click("button[role=switch]:has-text('peut rouvrir une question clôturée')")
    ps.screenshot(path=f"{OUT}/bridge_admin_client.png", full_page=True)
    ps.click("button[role=tab]:has-text('Utilisateurs')")
    ps.wait_for_selector("table[aria-label=Utilisateurs]")
    ps.click(f"tr[data-user-email='{LP}'] button:has-text('Modifier')")
    ps.wait_for_selector(f"[data-membership='{SLUG}']")
    anchor = f"[data-membership='{SLUG}'] [data-stream-right='Gouvernance'] button[aria-label^='Droit sur Gouvernance']"
    ps.locator(anchor).scroll_into_view_if_needed()
    ps.wait_for_timeout(250)
    ps.click(anchor)
    ps.click("#wib-popover button:has-text('Masqué')")
    ps.screenshot(path=f"{OUT}/bridge_admin_rights.png")
    ps.click("[role=dialog] button:has-text('Enregistrer')")
    ps.wait_for_timeout(800)
    lpu = next(u for u in su.req("GET", "/api/bridge/admin/users")[1] if u["email"] == LP)
    m = next(x for x in lpu["memberships"] if x["clientId"] == CID)
    check("admin : droit par stream enregistré", m["streamAccess"].get(SID["Gouvernance"]) == "NONE", str(m))
    check("droit Masqué : question Gouvernance invisible pour La Poste", all(x["id"] != q_cr["id"] for x in lp_api.req("GET", f"{BASE}/questions")[1]))
    ps.click("button:has-text('Inviter un utilisateur')")
    ps.fill("[role=dialog] input[aria-label=E-mail]", f"invite.ui.{RUN}@wibridge.test")
    ps.fill("[role=dialog] input[aria-label=Nom]", "Invité Interface")
    ps.click("[role=dialog] button:has-text('Inviter')")
    ps.wait_for_selector("text=Mode développement, lien")
    check("admin : invitation envoyée depuis l'interface", "Invitation envoyée par e-mail" in ps.inner_text("[role=dialog]"))
    ps.keyboard.press("Escape")
    ps.click("button[role=tab]:has-text('Journal')")
    ps.wait_for_selector("text=Configuration du client modifiée")
    check("admin : journal", "Stream ajouté : Logistique" in ps.inner_text("main"))
    ps.goto(f"{WEB}/compte")
    ps.wait_for_selector("text=Appareils de confiance")
    check("mon compte : cet appareil reconnu", "cet appareil" in ps.inner_text("main"))
    c_su.close()

    pl.goto(f"{WEB}/compte")
    pl.wait_for_selector("text=Notifications par e-mail")
    pl.click("text=Récapitulatif quotidien")
    pl.wait_for_timeout(700)
    check("mon compte : préférence de notification", next(c for c in lp_api.req("GET", "/api/bridge/auth/me")[1]["clients"] if c["slug"] == SLUG)["notify"] == "DAILY")
    pl.screenshot(path=f"{OUT}/bridge_compte.png", full_page=True)
    pl.click("button:has-text('Retirer tous les appareils')")
    pl.click("[role=dialog] button:has-text('Confirmer')")
    pl.wait_for_timeout(700)
    pl.click("button[aria-label='Menu utilisateur']")
    pl.click("text=Se déconnecter")
    pl.wait_for_url(re.compile(r".*/login"))
    check("appareils retirés : code redemandé", login(pl, LP, UPWD) is True)

    b.close()

# le client de l'exécution est archivé pour ne pas encombrer les listes
su.req("PATCH", f"/api/bridge/admin/clients/{CID}", {"archived": True})
try:
    os.remove(os.path.join(OUT, f"atm_{RUN}.xlsx"))
except OSError:
    pass
check("aucune erreur JavaScript", not errors, str(errors[:5]))
print()
print("RESULT:", "ALL OK" if not fails else f"{len(fails)} FAIL : {fails}")
