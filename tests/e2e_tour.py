"""Tour complet de WacMan : toutes les pages (ordinateur sombre, clair, mobile), rôle lecteur, comportements corrigés."""
import os as _os
ENV_FILE = _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "apps", "api", ".env.dev")
import sys, json, re, urllib.request, http.cookiejar, io
from playwright.sync_api import sync_playwright

BASE = "http://localhost:3000"
API = "http://localhost:4000"
OUT = sys.argv[1]
ENV = open(ENV_FILE).read()
PWD = ENV.split("BOOTSTRAP_ADMIN_PASSWORD=")[1].split("\n")[0].strip()
ACC = "la-poste-pstng"
VIEWER = ("lecteur.tour@wacman.test", "Lecteur-Tour-2026-ok!")
fails, errors = [], []


def check(name, cond, extra=""):
    print(("OK  " if cond else "FAIL"), name, extra)
    if not cond:
        fails.append(name)


cj = http.cookiejar.CookieJar()
op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))


def req(method, path, body=None):
    data = json.dumps(body).encode() if body is not None else None
    r = urllib.request.Request(API + path, data=data, method=method, headers={"Content-Type": "application/json"} if body is not None else {})
    try:
        with op.open(r) as res:
            t = res.read()
            return res.status, (json.loads(t) if t and "json" in res.headers.get("content-type", "") else t)
    except urllib.error.HTTPError as e:
        t = e.read().decode()
        try:
            return e.code, json.loads(t)
        except Exception:
            return e.code, t


s, r = req("POST", "/api/auth/login", {"email": "florent@omgt.fr", "password": PWD})
req("POST", "/api/auth/verify", {"challengeId": r["challengeId"], "code": r["devCode"]})
boot = req("GET", f"/api/accounts/{ACC}")[1]
types = boot["meetingTypes"]

# compte lecteur de test (local)
members = req("GET", f"/api/accounts/{ACC}/members")[1]
if not any(m.get("email") == VIEWER[0] for m in members):
    s, r = req("POST", f"/api/accounts/{ACC}/members", {"email": VIEWER[0], "name": "Lecteur Tour", "password": VIEWER[1], "role": "VIEWER"})
    print("création lecteur", s, r if s != 200 else "")

# exports : fichiers valides
from pptx import Presentation
import openpyxl
s, deck = req("GET", f"/api/accounts/{ACC}/export/deck.pptx?sections=cover,alerts,kanban,meetings,planning")
try:
    prs = Presentation(io.BytesIO(deck))
    txt = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    tbl = " ".join(c.text for sl in prs.slides for sh in sl.shapes if sh.has_table for row in sh.table.rows for c in row.cells)
    check("PPTX lisible", len(prs.slides) > 3, f"{len(prs.slides)} slides")
    check("PPTX sans balisage brut (**)", "**" not in txt + tbl)
except Exception as e:
    check("PPTX lisible", False, str(e))
s, x = req("GET", f"/api/accounts/{ACC}/export/cards.xlsx")
wb = openpyxl.load_workbook(io.BytesIO(x))
vals = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
check("Excel des cartes lisible et sans **", "**" not in vals and len(wb.worksheets) >= 1)
for t in types:
    s, x = req("GET", f"/api/accounts/{ACC}/export/meetings.xlsx?typeId={t['id']}")
    wb = openpyxl.load_workbook(io.BytesIO(x))
    check(f"Excel {t['name']} : au moins une feuille", len(wb.worksheets) >= 1)

pages = [("", "kanban"), ("/dashboard", "dashboard"), ("/risks", "risques"), ("/governance", "gouvernance"), ("/journal", "journal"), ("/finance", "finance"), ("/provisioning", "provisioning")]
pages += [(f"/meetings/{t['id']}", "seance_" + re.sub(r"\W+", "_", t["name"])) for t in types]
settings = ["general", "streams", "sprints", "lists", "meetings", "contacts", "members", "data"]


def login(page, email, pwd):
    page.goto(f"{BASE}/login")
    page.fill("input[type=email]", email)
    page.fill("input[type=password]", pwd)
    page.click("button:has-text('Continuer')")
    page.wait_for_selector("text=Mode développement")
    page.fill("input[autocomplete=one-time-code]", page.inner_text("text=Mode développement").split("code")[-1].strip())
    page.click("button:has-text('Se connecter')")
    page.wait_for_url(re.compile(r".*/(\?.*)?$|.*/a/.*"))


def watch(page, tag):
    page.on("pageerror", lambda e: errors.append(f"{tag} pageerror: {e}"))
    page.on("console", lambda m: m.type == "error" and "ERR_TUNNEL" not in m.text and "status of 403" not in m.text and errors.append(f"{tag} console: {m.text}"))
    page.on("response", lambda r: r.status >= 500 and errors.append(f"{tag} HTTP {r.status} {r.url}"))


with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(viewport={"width": 1440, "height": 950}, permissions=["clipboard-read", "clipboard-write"])
    page = ctx.new_page()
    watch(page, "admin")
    login(page, "florent@omgt.fr", PWD)

    # ---------------------------------------------------------------- toutes les pages, thème sombre
    for path, name in pages + [("/settings?tab=" + t, "param_" + t) for t in settings]:
        page.goto(f"{BASE}/a/{ACC}{path}")
        page.wait_for_timeout(1300)
        bad = page.locator("text=Une erreur").count() + page.locator("text=Compte inaccessible").count()
        check(f"page {name} s'affiche", bad == 0)
        page.screenshot(path=f"{OUT}/tour_{name}.png")
    for t in settings[1:]:
        pass
    page.goto(f"{BASE}/admin")
    page.wait_for_timeout(1200)
    check("page administration", page.locator("text=Une erreur").count() == 0)
    page.goto(f"{BASE}/")
    page.wait_for_timeout(800)
    check("liste des comptes", page.locator(f"a[href='/a/{ACC}']").count() >= 1)

    # ---------------------------------------------------------------- kanban : clavier, fenêtres empilées, archivage
    page.goto(f"{BASE}/a/{ACC}")
    page.click("button:has-text('Tous')")
    page.wait_for_timeout(800)
    first = page.locator("[aria-roledescription^=carte]").first
    first.focus()
    page.keyboard.press("Enter")
    page.wait_for_selector("[role=dialog]")
    check("carte ouverte au clavier (Entrée)", page.locator("[role=dialog]").count() >= 1)
    page.keyboard.press("Escape")
    page.wait_for_timeout(300)
    check("Échap ferme la carte", page.locator("[role=dialog]").count() == 0)

    # carte de test : création, date sans enregistrement à chaque chiffre, fenêtres empilées, archivage
    s, tc = req("POST", f"/api/accounts/{ACC}/e/card", {"title": "Carte tour " + "x" * 3 + "_tres_long_identifiant_sans_espace_pour_tester_le_retour_a_la_ligne_v12.xlsx", "statusId": [o for o in boot["options"] if o["kind"] == "CARD_STATUS"][0]["id"], "sprintId": [sp for sp in boot["sprints"] if sp["state"] == "CURRENT"][0]["id"]})
    page.reload()
    page.click("button:has-text('Tous')")
    page.wait_for_timeout(800)
    tile = page.locator("[aria-roledescription^=carte]", has_text="Carte tour").first
    box = tile.bounding_box()
    parent = tile.locator("xpath=..").bounding_box()
    check("titre très long : pas de débordement de la carte", box and parent and box["x"] + box["width"] <= parent["x"] + parent["width"] + 2, str((box, parent)))
    tile.click()
    page.wait_for_selector("[role=dialog]")
    patches = []
    page.on("request", lambda r: r.method == "PATCH" and "/e/card/" in r.url and patches.append(r.url))
    page.click("[role=dialog] button[aria-label='Ajouter : Échéance']")
    inp = page.locator("#wacman-popover input[type=date]")
    inp.fill("2026-11-14")
    page.wait_for_timeout(200)
    inp.fill("2026-11-15")
    page.wait_for_timeout(500)
    check("date : aucun enregistrement pendant la saisie", len(patches) == 0, str(len(patches)))
    page.click("#wacman-popover button:has-text('OK')")
    page.wait_for_timeout(800)
    tcr = req("GET", f"/api/accounts/{ACC}/e/card/{tc['id']}")[1]
    check("date : un seul enregistrement à la validation", len(patches) == 1 and tcr["dueDate"] == "2026-11-15", str((len(patches), tcr["dueDate"])))
    # fenêtres empilées : confirmation de suppression puis Échap
    page.click("[role=dialog] summary:has-text('Détails')")
    page.click("[role=dialog] button:has-text('Supprimer')")
    page.wait_for_timeout(300)
    n_dialogs = page.locator("[role=dialog]").count()
    page.keyboard.press("Escape")
    page.wait_for_timeout(300)
    check("Échap ne ferme que la confirmation", n_dialogs == 2 and page.locator("[role=dialog]").count() == 1, f"{n_dialogs} puis {page.locator('[role=dialog]').count()}")
    # barre de mise en forme : puce sur une ligne vide, Ctrl+K sans recherche
    page.locator("[role=dialog]").locator("xpath=.//*[normalize-space(text())='Prochaines étapes']/following::*[@role='button'][1]").first.click()
    ta = page.locator("[role=dialog] textarea").first
    ta.focus()
    page.click("[role=dialog] button[aria-label^='Liste à puces']")
    check("puce ajoutée sur une ligne vide", ta.input_value().startswith("- "), repr(ta.input_value()))
    page.keyboard.type("voir doc")
    page.keyboard.down("Shift")
    for _ in range(3):
        page.keyboard.press("ArrowLeft")
    page.keyboard.up("Shift")
    page.keyboard.press("Control+k")
    page.wait_for_timeout(300)
    check("Ctrl+K dans un texte : lien inséré, pas de recherche", "[doc](https://)" in ta.input_value() and page.locator("input[placeholder*='Rechercher']").count() == 0, repr(ta.input_value()))
    page.keyboard.press("Escape")
    page.wait_for_timeout(200)
    # archivage : la carte quitte le tableau
    page.click("[role=dialog] button:has-text('Archiver')")
    page.wait_for_timeout(600)
    if page.locator("[role=dialog]").count():
        page.keyboard.press("Escape")
    page.wait_for_timeout(500)
    check("carte archivée retirée du tableau", page.locator("[aria-roledescription^=carte]", has_text="Carte tour").count() == 0)
    req("DELETE", f"/api/accounts/{ACC}/e/card/{tc['id']}")

    # onglet « Sans sprint » : une carte créée là reste sans sprint
    page.click("button:has-text('Sans sprint')")
    page.click("button:has-text('Nouvelle carte')")
    page.fill("[role=dialog] input", "Carte sans sprint tour")
    page.click("[role=dialog] button:has-text('Créer la carte')")
    page.wait_for_timeout(1000)
    made = [c for c in req("GET", f"/api/accounts/{ACC}/cards")[1] if c["title"] == "Carte sans sprint tour"]
    check("carte créée depuis « Sans sprint » : sans sprint", len(made) == 1 and made[0]["sprintId"] is None, str([m["sprintId"] for m in made]))
    for m in made:
        req("DELETE", f"/api/accounts/{ACC}/e/card/{m['id']}")
    page.keyboard.press("Escape")

    # ---------------------------------------------------------------- thème clair
    page.evaluate("localStorage.setItem('wacman-theme','light')")
    for path, name in [("", "kanban"), ("/dashboard", "dashboard"), (f"/meetings/{types[0]['id']}", "pw")]:
        page.goto(f"{BASE}/a/{ACC}{path}")
        page.wait_for_timeout(1200)
        page.screenshot(path=f"{OUT}/tour_clair_{name}.png")
    page.evaluate("localStorage.removeItem('wacman-theme')")

    # ---------------------------------------------------------------- mobile
    mob = b.new_context(viewport={"width": 390, "height": 844}, storage_state=ctx.storage_state(), has_touch=True, is_mobile=True)
    mp = mob.new_page()
    watch(mp, "mobile")
    for path, name in pages + [("/settings", "parametres")]:
        mp.goto(f"{BASE}/a/{ACC}{path}")
        mp.wait_for_timeout(1300)
        w = mp.evaluate("document.documentElement.scrollWidth")
        check(f"mobile {name} sans défilement horizontal", w <= 392, f"scrollWidth={w}")
        mp.screenshot(path=f"{OUT}/tour_m_{name}.png")
    mob.close()

    # ---------------------------------------------------------------- redirection après connexion
    rc = b.new_context()
    rp = rc.new_page()
    rp.goto(f"{BASE}/login?next=//example.com/x")
    rp.fill("input[type=email]", "florent@omgt.fr")
    rp.fill("input[type=password]", PWD)
    rp.click("button:has-text('Continuer')")
    rp.wait_for_selector("text=Mode développement")
    rp.fill("input[autocomplete=one-time-code]", rp.inner_text("text=Mode développement").split("code")[-1].strip())
    rp.click("button:has-text('Se connecter')")
    rp.wait_for_timeout(2000)
    check("connexion : pas de redirection vers un autre site", rp.url.startswith(BASE), rp.url)
    rc.close()

    # ---------------------------------------------------------------- rôle lecteur
    vc = b.new_context(viewport={"width": 1440, "height": 950})
    vp = vc.new_page()
    watch(vp, "lecteur")
    login(vp, *VIEWER)
    for path, name in pages:
        vp.goto(f"{BASE}/a/{ACC}{path}")
        vp.wait_for_timeout(1200)
        check(f"lecteur {name} s'affiche", vp.locator("text=Une erreur").count() == 0)
    vp.goto(f"{BASE}/a/{ACC}")
    vp.wait_for_timeout(1000)
    check("lecteur : pas de « Nouvelle carte »", vp.locator("button:has-text('Nouvelle carte')").count() == 0)
    check("lecteur : pas de mode édition", vp.locator("button:has-text('Mode édition')").count() == 0)
    check("lecteur : pas de pastille d'ajout", vp.locator("button[aria-label^='Ajouter :']").count() == 0)
    vp.locator("[role=button]", has_text="#").first.click()
    vp.wait_for_timeout(800)
    check("lecteur : carte en lecture", vp.locator("[role=dialog]").count() == 1 and vp.locator("[role=dialog] button[aria-label^='Ajouter :']").count() == 0)
    vp.screenshot(path=f"{OUT}/tour_lecteur_carte.png")
    vc.close()
    b.close()

print("\nerrors:", "\n".join(errors[:25]) or "none")
print("RESULT:", "ALL OK" if not fails and not errors else f"{len(fails)} FAIL: {fails} / {len(errors)} errors")
